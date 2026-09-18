from __future__ import annotations

import hashlib
import json
import math
import os
import re
import sqlite3
import threading
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path

from flask import Flask, flash, g, redirect, render_template, request, send_file, session, url_for
from openpyxl import Workbook
from jinja2 import pass_context
from markupsafe import Markup
if __package__:
    from .localization import EMS_VI_TRANSLATIONS, StaticUiTranslation
else:
    from localization import EMS_VI_TRANSLATIONS, StaticUiTranslation


BASE_DIR = Path(__file__).resolve().parent
DB_PATH = Path(os.environ.get("ENERGY_DATABASE_PATH", BASE_DIR / "instance" / "app.db"))
POINT_SEED_PATH = BASE_DIR / "seed_data" / "utility_points.json"
POINT_MANIFEST_PATH = BASE_DIR / "seed_data" / "utility_points_manifest.json"
APP_VERSION = "0.7.5"
_DB_INIT_LOCK = threading.Lock()
_INITIALIZED_DB_PATH: Path | None = None

VI_TRANSLATIONS = {
    "能源資料填報系統": "Hệ thống khai báo dữ liệu năng lượng",
    "管理首頁": "Trang quản trị",
    "帳號權限管理": "Quản lý tài khoản và quyền",
    "能源/區域維護": "Quản lý năng lượng/khu vực",
    "填報查詢": "Tra cứu báo cáo",
    "資料查詢與匯出": "Tra cứu và xuất dữ liệu",
    "能源填報": "Khai báo năng lượng",
    "登出": "Đăng xuất",
    "登入": "Đăng nhập",
    "帳號": "Tài khoản",
    "密碼": "Mật khẩu",
    "管理者預設帳號": "Tài khoản quản trị mặc định",
    "能源資料輸入": "Nhập dữ liệu năng lượng",
    "填報月份": "Tháng khai báo",
    "每月填報上個月份資料；可選擇更早月份補填或修正。": "Mỗi tháng khai báo dữ liệu của tháng trước; có thể chọn tháng cũ hơn để bổ sung hoặc chỉnh sửa.",
    "填報月份只能選擇上個月或更早月份。": "Chỉ có thể chọn tháng trước hoặc tháng cũ hơn để khai báo.",
    "切換月份": "Đổi tháng",
    "填報資料": "Dữ liệu khai báo",
    "此月份資料已於": "Dữ liệu tháng này đã được gửi lúc",
    "送出。若需修正，可編輯後重新送出。": "Nếu cần sửa, có thể chỉnh sửa và gửi lại.",
    "請填寫所選月份資料。按下送出後，資料會交給管理者查詢；之後仍可修正並重新送出。": "Vui lòng nhập dữ liệu của tháng đã chọn. Sau khi gửi, quản trị viên có thể tra cứu; bạn vẫn có thể sửa và gửi lại.",
    "Excel 複製貼上操作": "Cách sao chép và dán từ Excel",
    "可以直接從 Excel 複製多格資料，並一次貼到下方填報表格。": "Bạn có thể sao chép nhiều ô từ Excel và dán một lần vào bảng khai báo bên dưới.",
    "在 Excel 依照填報表格的順序選取資料範圍。": "Trong Excel, chọn vùng dữ liệu theo đúng thứ tự của bảng khai báo.",
    "按 Ctrl+C 複製選取的 Excel 資料。": "Nhấn Ctrl+C để sao chép dữ liệu đã chọn trong Excel.",
    "回到本頁，點擊要開始填入的第一個欄位。": "Quay lại trang này và nhấp vào ô đầu tiên muốn điền.",
    "按 Ctrl+V，資料會依照 Excel 的列與欄自動填入。": "Nhấn Ctrl+V; dữ liệu sẽ tự động điền theo hàng và cột trong Excel.",
    "確認綠色欄位與提示訊息；若有紅色錯誤欄位，請修正後再送出。": "Kiểm tra các ô màu xanh và thông báo; nếu có ô lỗi màu đỏ, hãy sửa trước khi gửi.",
    "注意：請不要複製 Excel 的標題列，只需複製數量與金額資料。": "Lưu ý: Không sao chép hàng tiêu đề trong Excel; chỉ sao chép dữ liệu số lượng và số tiền.",
    "管制值管理": "Quản lý giới hạn kiểm soát",
    "依歷史平均值設定各能源種類與單位的管制範圍。超出範圍時會警告，但不會阻止送出。": "Thiết lập phạm vi kiểm soát theo giá trị trung bình lịch sử cho từng loại năng lượng và đơn vị. Hệ thống sẽ cảnh báo khi vượt phạm vi nhưng vẫn cho phép gửi.",
    "重新依歷史平均值計算全部項目（±10%）": "Tính lại tất cả theo trung bình lịch sử (±10%)",
    "數量管制": "Kiểm soát số lượng",
    "金額管制": "Kiểm soát số tiền",
    "歷史平均": "Trung bình lịch sử",
    "下限": "Giới hạn dưới",
    "上限": "Giới hạn trên",
    "啟用": "Bật",
    "儲存管制值": "Lưu giới hạn kiểm soát",
    "尚無歷史資料": "Chưa có dữ liệu lịch sử",
    "輸入值超出管制範圍": "Giá trị nhập vượt phạm vi kiểm soát",
    "筆資料超出管制範圍，仍可送出。": "giá trị vượt phạm vi kiểm soát; vẫn có thể gửi.",
    "管制值已更新。": "Đã cập nhật giới hạn kiểm soát.",
    "管制值已依歷史平均值重新計算。": "Đã tính lại giới hạn theo trung bình lịch sử.",
    "管制值格式不正確，請確認上下限。": "Định dạng giới hạn không hợp lệ; vui lòng kiểm tra giới hạn dưới và trên.",
    "下限不可大於上限。": "Giới hạn dưới không được lớn hơn giới hạn trên.",
    "確定要送出": "Bạn có chắc muốn gửi dữ liệu tháng",
    "的填報資料嗎？": "không?",
    "送出後仍可修改並重新送出。": "Sau khi gửi, bạn vẫn có thể chỉnh sửa và gửi lại.",
    "目前有資料超出管制範圍，確定仍要送出嗎？": "Hiện có dữ liệu vượt phạm vi kiểm soát. Bạn vẫn muốn gửi?",
    "超出管制值明細：": "Chi tiết vượt giới hạn kiểm soát:",
    "輸入值：": "Giá trị nhập:",
    "管制範圍：": "Phạm vi kiểm soát:",
    "點擊可移至該欄位": "Nhấp để chuyển đến ô này",
    "目前尚未被指派可填寫的能源種類或單位，請洽管理者設定。": "Bạn chưa được cấp quyền cho loại năng lượng hoặc đơn vị. Vui lòng liên hệ quản trị viên.",
    "能源種類 / 單位": "Loại năng lượng / Đơn vị",
    "數量": "Số lượng",
    "金額(未稅)": "Số tiền (chưa thuế)",
    "重新送出填報資料": "Gửi lại dữ liệu",
    "送出填報資料": "Gửi dữ liệu",
    "請輸入有效數字": "Vui lòng nhập số hợp lệ",
    "貼上的內容不是有效數字。": "Nội dung dán không phải là số hợp lệ.",
    "不是有效數字": "không phải là số hợp lệ",
    "已填入": "Đã điền",
    "格超出表格範圍": "ô vượt ngoài phạm vi bảng",
    "格": "ô",
    "Excel 貼上成功，共填入": "Dán từ Excel thành công, tổng cộng",
    "帳號與權限管理": "Quản lý tài khoản và quyền",
    "新增填寫員或訪客": "Thêm nhân viên nhập liệu hoặc khách",
    "顯示名稱": "Tên hiển thị",
    "帳號權限": "Quyền tài khoản",
    "填寫員": "Nhân viên nhập liệu",
    "訪客（僅查看及匯出）": "Khách (chỉ xem và xuất)",
    "建立": "Tạo",
    "訪客（唯讀）": "Khách (chỉ đọc)",
    "Y 軸：可填寫能源種類": "Trục Y: Loại năng lượng được phép nhập",
    "X 軸：可填寫單位": "Trục X: Đơn vị được phép nhập",
    "儲存權限": "Lưu quyền",
    "此帳號可查看所有年份資料並匯出 Excel，不可填報或修改資料。": "Tài khoản này có thể xem dữ liệu mọi năm và xuất Excel, nhưng không thể nhập hoặc sửa dữ liệu.",
    "重設密碼": "Đặt lại mật khẩu",
    "輸入新密碼": "Nhập mật khẩu mới",
    "尚未建立填寫員。": "Chưa có tài khoản.",
    "指定月份（優先）": "Chọn tháng (ưu tiên)",
    "指定月份": "Chọn tháng",
    "起始月份": "Tháng bắt đầu",
    "結束月份": "Tháng kết thúc",
    "區間包含起始與結束月份，可跨年度。查詢優先順序：月份區間、指定月份、年份；全部留空則查全部資料。": "Khoảng bao gồm tháng bắt đầu và kết thúc, có thể qua nhiều năm. Ưu tiên: khoảng tháng, tháng cụ thể, năm; để trống tất cả để xem toàn bộ dữ liệu.",
    "請同時填寫起始月份與結束月份。": "Vui lòng nhập cả tháng bắt đầu và tháng kết thúc.",
    "月份區間格式不正確。": "Định dạng khoảng tháng không hợp lệ.",
    "起始月份不可晚於結束月份。": "Tháng bắt đầu không được sau tháng kết thúc.",
    "或選擇年份": "Hoặc chọn năm",
    "全部年份": "Tất cả các năm",
    "查詢": "Tra cứu",
    "全部資料": "Tất cả dữ liệu",
    "匯出 Excel": "Xuất Excel",
    "月份": "Tháng",
    "能源種類": "Loại năng lượng",
    "能源單位": "Đơn vị năng lượng",
    "單位": "Đơn vị",
    "區域": "Khu vực",
    "送出時間": "Thời gian gửi",
    "更新時間": "Thời gian cập nhật",
    "查詢範圍內尚無填報資料。": "Không có dữ liệu trong phạm vi tìm kiếm.",
    "新增能源種類": "Thêm loại năng lượng",
    "例如：水(in-自來水)": "Ví dụ: Nước (nước máy vào)",
    "例如：公噸": "Ví dụ: tấn",
    "新增區域": "Thêm khu vực",
    "區域名稱": "Tên khu vực",
    "例如：Poly 51": "Ví dụ: Poly 51",
    "目前能源種類": "Các loại năng lượng hiện tại",
    "排序": "Thứ tự",
    "目前區域": "Các khu vực hiện tại",
    "能源分析看板": "Bảng phân tích năng lượng",
    "掌握各月份能源費用、填報進度與主要成本分布。": "Theo dõi chi phí năng lượng, tiến độ khai báo và phân bổ chi phí chính theo tháng.",
    "分析月份": "Tháng phân tích",
    "更新": "Cập nhật",
    "能源總金額（未稅）": "Tổng chi phí năng lượng (chưa thuế)",
    "筆有效填報資料": "bản ghi hợp lệ",
    "較": "So với",
    "增減": "thay đổi",
    "上月": "Tháng trước",
    "上月沒有可比較資料": "Không có dữ liệu tháng trước để so sánh",
    "已填報人數": "Số người đã khai báo",
    "依有效填寫員帳號計算": "Tính theo tài khoản nhập liệu đang hoạt động",
    "尚未填報人數": "Số người chưa khai báo",
    "仍需追蹤填報進度": "Cần tiếp tục theo dõi tiến độ",
    "本月皆已完成送出": "Tất cả đã gửi trong tháng này",
    "每月能源費用趨勢": "Xu hướng chi phí năng lượng hàng tháng",
    "最近 12 個有資料的月份": "12 tháng gần nhất có dữ liệu",
    "目前尚無趨勢資料。": "Hiện chưa có dữ liệu xu hướng.",
    "能源費用占比": "Tỷ trọng chi phí năng lượng",
    "依未稅金額排序": "Sắp xếp theo số tiền chưa thuế",
    "此月份尚無能源金額資料。": "Tháng này chưa có dữ liệu chi phí năng lượng.",
    "各區域費用排名": "Xếp hạng chi phí theo khu vực",
    "含占當月總費用比例": "Bao gồm tỷ lệ trong tổng chi phí tháng",
    "目前尚未設定區域。": "Hiện chưa thiết lập khu vực.",
    "請先登入。": "Vui lòng đăng nhập trước.",
    "此功能限管理者使用。": "Chức năng này chỉ dành cho quản trị viên.",
    "此功能限管理者或訪客使用。": "Chức năng này chỉ dành cho quản trị viên hoặc khách.",
    "登入成功。": "Đăng nhập thành công.",
    "帳號或密碼錯誤。": "Tài khoản hoặc mật khẩu không đúng.",
    "已登出。": "Đã đăng xuất.",
    "分析月份格式不正確。": "Định dạng tháng phân tích không hợp lệ.",
    "請輸入帳號與密碼。": "Vui lòng nhập tài khoản và mật khẩu.",
    "填寫員已建立。": "Đã tạo tài khoản.",
    "帳號已存在。": "Tài khoản đã tồn tại.",
    "找不到填寫員。": "Không tìm thấy tài khoản.",
    "填寫權限已更新。": "Đã cập nhật quyền nhập liệu.",
    "請輸入新密碼。": "Vui lòng nhập mật khẩu mới.",
    "密碼已重設。": "Đã đặt lại mật khẩu.",
    "請輸入能源種類與單位。": "Vui lòng nhập loại năng lượng và đơn vị.",
    "能源種類已新增。": "Đã thêm loại năng lượng.",
    "能源種類已存在。": "Loại năng lượng đã tồn tại.",
    "請輸入區域名稱。": "Vui lòng nhập tên khu vực.",
    "區域已新增。": "Đã thêm khu vực.",
    "區域已存在。": "Khu vực đã tồn tại.",
    "填報月份格式不正確。": "Định dạng tháng khai báo không hợp lệ.",
    "能源資料已送出。若需修正，可再次編輯後重新送出。": "Dữ liệu năng lượng đã được gửi. Nếu cần sửa, hãy chỉnh sửa và gửi lại.",
    "能源填報資料": "Dữ liệu khai báo năng lượng",
    "年": "Năm",
    "備份資料庫": "Sao lưu cơ sở dữ liệu",
    "系統日誌": "Nhật ký hệ thống",
    "時間": "Thời gian",
    "人員": "Nhân viên",
    "動作": "Hành động",
    "詳細內容": "Chi tiết",
    "重新計算偏離值範圍：": "Tính lại phạm vi độ lệch:",
    "重新計算全部項目": "Tính lại tất cả các mục",
    "資料庫備份成功。": "Sao lưu cơ sở dữ liệu thành công.",
}


VI_TRANSLATIONS.update(EMS_VI_TRANSLATIONS)
_UI_PATTERN = re.compile("|".join(re.escape(key) for key in sorted(VI_TRANSLATIONS, key=len, reverse=True)))

app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("ENERGY_SECRET_KEY", "dev-secret-change-me")


def translate(text: str) -> str:
    if session.get("language", "zh-TW") == "vi":
        return VI_TRANSLATIONS.get(text, text)
    return text


@app.context_processor
def inject_app_version() -> dict[str, str]:
    return {
        "app_version": APP_VERSION,
        "current_language": session.get("language", "zh-TW"),
    }


def translate_ui_text(text: str) -> str:
    if session.get("language") == "vi":
        return _UI_PATTERN.sub(lambda match: VI_TRANSLATIONS[match.group()], str(text))
    return str(text)


@pass_context
def translate_static_ui(context, text: str) -> Markup:
    # This filter is injected only for literal, trusted template source.
    return Markup(translate_ui_text(text))


app.jinja_env.filters["ui_static"] = translate_static_ui
app.jinja_env.filters["ui"] = pass_context(lambda context, text: translate_ui_text(text))
app.jinja_env.filters["label"] = pass_context(lambda context, text: translate(text))
app.jinja_env.add_extension(StaticUiTranslation)


@app.route("/language/<language>", methods=["POST"])
def set_language(language: str):
    if language in {"zh-TW", "vi"}:
        session["language"] = language
    next_url = request.form.get("next", "")
    if not next_url.startswith("/") or next_url.startswith("//"):
        next_url = url_for("index")
    return redirect(next_url)


DEFAULT_ENERGY_TYPES = [
    ("產量(T)", "T"),
    ("水(in-自來水)", "公噸"),
    ("水(out-排放水)", "公噸"),
    ("水(in-湖水)", "公噸"),
    ("氮氣(液氮)", "公噸"),
    ("電(台電、供電局)", "度"),
    ("電(尖峰)", "度"),
    ("電(離峰)", "度"),
    ("電(正常)", "度"),
    ("蒸汽(自產-北廠熱煤鍋爐)", "公噸"),
    ("蒸汽(自產-南廠蒸汽鍋爐)", "公噸"),
    ("煤碳(北廠熱煤鍋爐)", "公噸"),
    ("煤碳(南廠蒸汽鍋爐)", "公噸"),
    ("煤碳-蒸汽(南廠蒸汽鍋爐)", "公噸"),
    ("煤碳-熱媒(北廠熱煤鍋爐)", "公噸"),
    ("木屑粒(北廠熱煤鍋爐)", "公噸"),
    ("天然氣(北廠熱媒鍋爐)", "Sm3"),
    ("稻殼粒(南廠蒸汽鍋爐)", "公噸"),
]

DEFAULT_AREAS = [
    "Poly 51",
    "SSP 51",
    "SSP Rpet",
    "Poly 52",
    "FL 51",
    "POY",
    "DTY",
    "POY造粒",
]


def current_report_month() -> str:
    bangkok_time = datetime.now(timezone(timedelta(hours=7)))
    return bangkok_time.strftime("%Y-%m")


def latest_report_month(today: date | None = None) -> str:
    reference = today or datetime.now(timezone(timedelta(hours=7))).date()
    first_day_of_month = reference.replace(day=1)
    return (first_day_of_month - timedelta(days=1)).strftime("%Y-%m")


def valid_report_month(value: str) -> bool:
    try:
        date.fromisoformat(f"{value}-01")
    except ValueError:
        return False
    return len(value) == 7


def previous_report_month(value: str) -> str:
    year, month = (int(part) for part in value.split("-"))
    if month == 1:
        return f"{year - 1}-12"
    return f"{year}-{month - 1:02d}"


def get_db() -> sqlite3.Connection:
    if "db" not in g:
        DB_PATH.parent.mkdir(parents=True, exist_ok=True)
        g.db = sqlite3.connect(DB_PATH, timeout=10)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
        g.db.execute("PRAGMA busy_timeout = 10000")
    return g.db


@app.teardown_appcontext
def close_db(error: BaseException | None) -> None:
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db() -> None:
    db = get_db()
    area_permissions_exist = db.execute(
        """
        SELECT 1 FROM sqlite_master
        WHERE type = 'table' AND name = 'user_area_permissions'
        """
    ).fetchone() is not None
    db.executescript(
        """
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL UNIQUE,
            password TEXT NOT NULL,
            display_name TEXT NOT NULL,
            role TEXT NOT NULL CHECK(role IN ('admin', 'filler')),
            viewer INTEGER NOT NULL DEFAULT 0,
            active INTEGER NOT NULL DEFAULT 1,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );

        CREATE TABLE IF NOT EXISTS energy_types (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL UNIQUE,
            unit TEXT NOT NULL,
            active INTEGER NOT NULL DEFAULT 1,
            display_order INTEGER NOT NULL DEFAULT 0
        );

        CREATE TABLE IF NOT EXISTS areas (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL UNIQUE,
            active INTEGER NOT NULL DEFAULT 1,
            display_order INTEGER NOT NULL DEFAULT 0
        );

        CREATE TABLE IF NOT EXISTS user_energy_permissions (
            user_id INTEGER NOT NULL,
            energy_type_id INTEGER NOT NULL,
            PRIMARY KEY (user_id, energy_type_id),
            FOREIGN KEY (user_id) REFERENCES users(id),
            FOREIGN KEY (energy_type_id) REFERENCES energy_types(id)
        );

        CREATE TABLE IF NOT EXISTS user_area_permissions (
            user_id INTEGER NOT NULL,
            area_id INTEGER NOT NULL,
            PRIMARY KEY (user_id, area_id),
            FOREIGN KEY (user_id) REFERENCES users(id),
            FOREIGN KEY (area_id) REFERENCES areas(id)
        );

        CREATE TABLE IF NOT EXISTS energy_entries (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            report_month TEXT NOT NULL,
            user_id INTEGER NOT NULL,
            energy_type_id INTEGER NOT NULL,
            area_id INTEGER NOT NULL,
            quantity REAL,
            amount REAL,
            updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(report_month, user_id, energy_type_id, area_id),
            FOREIGN KEY (user_id) REFERENCES users(id),
            FOREIGN KEY (energy_type_id) REFERENCES energy_types(id),
            FOREIGN KEY (area_id) REFERENCES areas(id)
        );

        CREATE TABLE IF NOT EXISTS report_submissions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            report_month TEXT NOT NULL,
            user_id INTEGER NOT NULL,
            status TEXT NOT NULL DEFAULT 'submitted',
            submitted_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(report_month, user_id),
            FOREIGN KEY (user_id) REFERENCES users(id)
        );

        CREATE TABLE IF NOT EXISTS control_limits (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            energy_type_id INTEGER NOT NULL,
            area_id INTEGER NOT NULL,
            quantity_enabled INTEGER NOT NULL DEFAULT 0,
            quantity_average REAL,
            quantity_min REAL,
            quantity_max REAL,
            amount_enabled INTEGER NOT NULL DEFAULT 0,
            amount_average REAL,
            amount_min REAL,
            amount_max REAL,
            updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(energy_type_id, area_id),
            FOREIGN KEY (energy_type_id) REFERENCES energy_types(id),
            FOREIGN KEY (area_id) REFERENCES areas(id)
        );

        CREATE TABLE IF NOT EXISTS audit_logs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER,
            username TEXT,
            action TEXT NOT NULL,
            details TEXT,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (user_id) REFERENCES users(id)
        );

        CREATE TABLE IF NOT EXISTS device_sources (
            external_id INTEGER PRIMARY KEY,
            name TEXT NOT NULL,
            api_url TEXT,
            source_is_active INTEGER NOT NULL DEFAULT 1,
            last_synced_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );

        CREATE TABLE IF NOT EXISTS device_data_points (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            external_id INTEGER NOT NULL,
            data_source_id INTEGER NOT NULL,
            point_name TEXT,
            json_path TEXT,
            description TEXT,
            source_is_active INTEGER NOT NULL DEFAULT 1,
            point_type TEXT,
            fetch_interval_minutes INTEGER,
            source_present INTEGER NOT NULL DEFAULT 1,
            source_checksum TEXT,
            last_synced_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(data_source_id, external_id),
            FOREIGN KEY (data_source_id) REFERENCES device_sources(external_id)
        );

        CREATE TABLE IF NOT EXISTS ems_energy_categories (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL UNIQUE,
            display_order INTEGER NOT NULL DEFAULT 0,
            active INTEGER NOT NULL DEFAULT 1
        );

        CREATE TABLE IF NOT EXISTS ems_measurement_types (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL UNIQUE,
            display_order INTEGER NOT NULL DEFAULT 0,
            active INTEGER NOT NULL DEFAULT 1
        );

        CREATE TABLE IF NOT EXISTS ems_units (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL UNIQUE,
            display_order INTEGER NOT NULL DEFAULT 0,
            active INTEGER NOT NULL DEFAULT 1
        );

        CREATE TABLE IF NOT EXISTS ems_locations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL UNIQUE,
            display_order INTEGER NOT NULL DEFAULT 0,
            active INTEGER NOT NULL DEFAULT 1
        );

        CREATE TABLE IF NOT EXISTS ems_point_configs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            point_id INTEGER NOT NULL UNIQUE,
            monitor_enabled INTEGER NOT NULL DEFAULT 1,
            lower_limit REAL,
            upper_limit REAL,
            display_name TEXT,
            energy_category_id INTEGER,
            measurement_type_id INTEGER,
            unit_id INTEGER,
            location_id INTEGER,
            equipment_name TEXT,
            value_type TEXT,
            aggregation_method TEXT,
            flow_direction TEXT,
            alarm_severity TEXT NOT NULL DEFAULT 'warning',
            alarm_delay_minutes INTEGER NOT NULL DEFAULT 5,
            alarm_value_mode TEXT NOT NULL DEFAULT 'raw',
            alarm_consecutive_samples INTEGER NOT NULL DEFAULT 1,
            alarm_deadband REAL NOT NULL DEFAULT 0,
            notes TEXT,
            updated_by INTEGER,
            updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (point_id) REFERENCES device_data_points(id),
            FOREIGN KEY (energy_category_id) REFERENCES ems_energy_categories(id),
            FOREIGN KEY (measurement_type_id) REFERENCES ems_measurement_types(id),
            FOREIGN KEY (unit_id) REFERENCES ems_units(id),
            FOREIGN KEY (location_id) REFERENCES ems_locations(id),
            FOREIGN KEY (updated_by) REFERENCES users(id)
        );

        CREATE TABLE IF NOT EXISTS point_catalog_versions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            data_source_id INTEGER NOT NULL,
            version TEXT NOT NULL,
            checksum TEXT NOT NULL,
            record_count INTEGER NOT NULL,
            added_count INTEGER NOT NULL DEFAULT 0,
            updated_count INTEGER NOT NULL DEFAULT 0,
            missing_count INTEGER NOT NULL DEFAULT 0,
            status TEXT NOT NULL,
            details TEXT,
            imported_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );

        CREATE INDEX IF NOT EXISTS idx_device_points_source
            ON device_data_points(data_source_id, source_present, source_is_active);
        CREATE INDEX IF NOT EXISTS idx_ems_configs_monitor
            ON ems_point_configs(monitor_enabled);
        CREATE INDEX IF NOT EXISTS idx_point_catalog_source
            ON point_catalog_versions(data_source_id, imported_at);
        """
    )
    if "viewer" not in {
        row["name"] for row in db.execute("PRAGMA table_info(users)").fetchall()
    }:
        db.execute("ALTER TABLE users ADD COLUMN viewer INTEGER NOT NULL DEFAULT 0")
    ems_config_columns = {
        row["name"] for row in db.execute("PRAGMA table_info(ems_point_configs)").fetchall()
    }
    if "alarm_value_mode" not in ems_config_columns:
        db.execute(
            "ALTER TABLE ems_point_configs ADD COLUMN alarm_value_mode TEXT NOT NULL DEFAULT 'raw'"
        )
    if "alarm_consecutive_samples" not in ems_config_columns:
        db.execute(
            "ALTER TABLE ems_point_configs "
            "ADD COLUMN alarm_consecutive_samples INTEGER NOT NULL DEFAULT 1"
        )
    seed_defaults(db)
    seed_control_limits(db)
    seed_ems_lookups(db)
    if not area_permissions_exist:
        # Existing fillers used to have access to every area. Preserve that
        # behavior when introducing X-axis permissions.
        db.execute(
            """
            INSERT OR IGNORE INTO user_area_permissions (user_id, area_id)
            SELECT u.id, a.id
            FROM users u
            CROSS JOIN areas a
            WHERE u.role = 'filler' AND a.active = 1
            """
        )
    db.commit()
    sync_bundled_device_points(db)


def seed_control_limits(db: sqlite3.Connection, replace: bool = False, percentage: float = 0.1) -> None:
    if replace:
        db.execute("DELETE FROM control_limits")
    db.execute(
        """
        WITH averages AS (
            SELECT
                energy_type_id,
                area_id,
                AVG(quantity) AS quantity_average,
                AVG(amount) AS amount_average
            FROM energy_entries
            GROUP BY energy_type_id, area_id
        )
        INSERT OR IGNORE INTO control_limits (
            energy_type_id, area_id,
            quantity_enabled, quantity_average, quantity_min, quantity_max,
            amount_enabled, amount_average, amount_min, amount_max
        )
        SELECT
            et.id,
            a.id,
            CASE WHEN av.quantity_average IS NULL THEN 0 ELSE 1 END,
            av.quantity_average,
            av.quantity_average - ABS(av.quantity_average) * ?,
            av.quantity_average + ABS(av.quantity_average) * ?,
            CASE WHEN av.amount_average IS NULL THEN 0 ELSE 1 END,
            av.amount_average,
            av.amount_average - ABS(av.amount_average) * ?,
            av.amount_average + ABS(av.amount_average) * ?
        FROM energy_types et
        CROSS JOIN areas a
        LEFT JOIN averages av
            ON av.energy_type_id = et.id
            AND av.area_id = a.id
        WHERE et.active = 1 AND a.active = 1
        """,
        (percentage, percentage, percentage, percentage),
    )


def log_action(action: str, details: str = None) -> None:
    try:
        db = get_db()
        user_id = session.get("user_id")
        username = session.get("username", "system")
        db.execute(
            "INSERT INTO audit_logs (user_id, username, action, details) VALUES (?, ?, ?, ?)",
            (user_id, username, action, details),
        )
        db.commit()
    except Exception:
        pass


def seed_defaults(db: sqlite3.Connection) -> None:
    db.execute(
        """
        INSERT OR IGNORE INTO users (username, password, display_name, role)
        VALUES ('admin', '1234', '系統管理者', 'admin')
        """
    )
    for index, (name, unit) in enumerate(DEFAULT_ENERGY_TYPES, start=1):
        db.execute(
            """
            INSERT OR IGNORE INTO energy_types (name, unit, display_order)
            VALUES (?, ?, ?)
            """,
            (name, unit, index),
        )
    for index, name in enumerate(DEFAULT_AREAS, start=1):
        db.execute(
            """
            INSERT OR IGNORE INTO areas (name, display_order)
            VALUES (?, ?)
            """,
            (name, index),
        )


def seed_ems_lookups(db: sqlite3.Connection) -> None:
    lookups = {
        "ems_energy_categories": [
            "電力", "用水", "廢水", "蒸汽", "天然氣", "壓縮空氣", "氮氣",
            "冰水", "冷卻水", "煤", "生質燃料", "產量", "環境指標", "其他",
        ],
        "ems_measurement_types": [
            "瞬時流量", "累積流量", "功率", "累積能源", "壓力", "溫度",
            "露點", "液位", "導電度", "COD", "電流", "電壓", "運轉狀態",
            "產量", "其他",
        ],
        "ems_units": [
            "m³/h", "m³", "Nm³/h", "Nm³", "bar", "kPa", "°C", "kW", "kWh",
            "A", "V", "mg/L", "公噸", "狀態", "%",
        ],
        "ems_locations": ["Utility", "POY", "DTY", "SSP", "SSP Rpet", "PSF", "WWT", "4ha"],
    }
    for table, names in lookups.items():
        db.executemany(
            f"INSERT OR IGNORE INTO {table} (name, display_order) VALUES (?, ?)",
            [(name, index) for index, name in enumerate(names, start=1)],
        )


def _point_metadata_checksum(point: dict) -> str:
    fields = {
        "external_id": point.get("external_id"),
        "data_source_id": point.get("data_source_id"),
        "point_name": point.get("point_name") or "",
        "json_path": point.get("json_path") or "",
        "description": point.get("description") or "",
        "source_is_active": bool(point.get("source_is_active")),
        "point_type": point.get("point_type") or "",
        "fetch_interval_minutes": point.get("fetch_interval_minutes"),
    }
    raw = json.dumps(fields, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def load_bundled_point_catalog() -> tuple[dict, dict] | None:
    if not POINT_SEED_PATH.exists() or not POINT_MANIFEST_PATH.exists():
        return None

    seed_bytes = POINT_SEED_PATH.read_bytes()
    manifest = json.loads(POINT_MANIFEST_PATH.read_text(encoding="utf-8"))
    actual_checksum = hashlib.sha256(seed_bytes).hexdigest()
    if manifest.get("sha256") != actual_checksum:
        raise ValueError("點位種子資料 checksum 不符")

    catalog = json.loads(seed_bytes.decode("utf-8"))
    points = catalog.get("points")
    source = catalog.get("source")
    if not isinstance(points, list) or not isinstance(source, dict):
        raise ValueError("點位種子資料格式不正確")
    if manifest.get("row_count") != len(points):
        raise ValueError("點位種子資料筆數與 manifest 不符")
    if manifest.get("data_source_id") != source.get("external_id"):
        raise ValueError("點位種子資料來源與 manifest 不符")

    external_ids = [point.get("external_id") for point in points]
    if any(not isinstance(value, int) for value in external_ids):
        raise ValueError("點位種子資料包含無效的外部 Id")
    if len(external_ids) != len(set(external_ids)):
        raise ValueError("點位種子資料包含重複的外部 Id")
    if any(point.get("data_source_id") != source.get("external_id") for point in points):
        raise ValueError("點位種子資料包含錯誤的 DataSourceId")
    return catalog, manifest


def sync_bundled_device_points(db: sqlite3.Connection) -> dict | None:
    try:
        loaded = load_bundled_point_catalog()
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        app.logger.exception("Cannot load bundled point catalog: %s", exc)
        return {"status": "failed", "details": str(exc)}
    if loaded is None:
        return None

    catalog, manifest = loaded
    source = catalog["source"]
    points = catalog["points"]
    data_source_id = int(source["external_id"])
    checksum = manifest["sha256"]
    latest = db.execute(
        """
        SELECT * FROM point_catalog_versions
        WHERE data_source_id = ? AND status = 'success'
        ORDER BY id DESC LIMIT 1
        """,
        (data_source_id,),
    ).fetchone()
    if latest and latest["checksum"] == checksum:
        return {"status": "current", "version": latest["version"], "record_count": latest["record_count"]}

    try:
        db.execute("BEGIN IMMEDIATE")
        latest = db.execute(
            """
            SELECT * FROM point_catalog_versions
            WHERE data_source_id = ? AND status = 'success'
            ORDER BY id DESC LIMIT 1
            """,
            (data_source_id,),
        ).fetchone()
        if latest and latest["checksum"] == checksum:
            db.commit()
            return {"status": "current", "version": latest["version"], "record_count": latest["record_count"]}

        db.execute(
            """
            INSERT INTO device_sources (
                external_id, name, api_url, source_is_active, last_synced_at
            ) VALUES (?, ?, ?, ?, CURRENT_TIMESTAMP)
            ON CONFLICT(external_id) DO UPDATE SET
                name = excluded.name,
                api_url = COALESCE(excluded.api_url, device_sources.api_url),
                source_is_active = excluded.source_is_active,
                last_synced_at = CURRENT_TIMESTAMP
            """,
            (
                data_source_id,
                source.get("name") or f"DataSource {data_source_id}",
                source.get("api_url"),
                bool(source.get("source_is_active", True)),
            ),
        )
        existing = {
            row["external_id"]: row
            for row in db.execute(
                "SELECT * FROM device_data_points WHERE data_source_id = ?",
                (data_source_id,),
            ).fetchall()
        }
        db.execute(
            "UPDATE device_data_points SET source_present = 0 WHERE data_source_id = ?",
            (data_source_id,),
        )
        added_count = 0
        updated_count = 0
        for point in points:
            external_id = int(point["external_id"])
            point_checksum = _point_metadata_checksum(point)
            old = existing.get(external_id)
            if old is None:
                added_count += 1
            elif old["source_checksum"] != point_checksum or not old["source_present"]:
                updated_count += 1
            db.execute(
                """
                INSERT INTO device_data_points (
                    external_id, data_source_id, point_name, json_path, description,
                    source_is_active, point_type, fetch_interval_minutes,
                    source_present, source_checksum, last_synced_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, 1, ?, CURRENT_TIMESTAMP)
                ON CONFLICT(data_source_id, external_id) DO UPDATE SET
                    point_name = excluded.point_name,
                    json_path = excluded.json_path,
                    description = excluded.description,
                    source_is_active = excluded.source_is_active,
                    point_type = excluded.point_type,
                    fetch_interval_minutes = excluded.fetch_interval_minutes,
                    source_present = 1,
                    source_checksum = excluded.source_checksum,
                    last_synced_at = CURRENT_TIMESTAMP
                """,
                (
                    external_id,
                    data_source_id,
                    point.get("point_name"),
                    point.get("json_path"),
                    point.get("description"),
                    bool(point.get("source_is_active", True)),
                    point.get("point_type"),
                    point.get("fetch_interval_minutes"),
                    point_checksum,
                ),
            )
            local_point_id = db.execute(
                "SELECT id FROM device_data_points WHERE data_source_id = ? AND external_id = ?",
                (data_source_id, external_id),
            ).fetchone()["id"]
            db.execute(
                "INSERT OR IGNORE INTO ems_point_configs (point_id, monitor_enabled) VALUES (?, 1)",
                (local_point_id,),
            )

        missing_count = db.execute(
            "SELECT COUNT(*) FROM device_data_points WHERE data_source_id = ? AND source_present = 0",
            (data_source_id,),
        ).fetchone()[0]
        details = f"Added {added_count}, updated {updated_count}, missing {missing_count}"
        db.execute(
            """
            INSERT INTO point_catalog_versions (
                data_source_id, version, checksum, record_count,
                added_count, updated_count, missing_count, status, details
            ) VALUES (?, ?, ?, ?, ?, ?, ?, 'success', ?)
            """,
            (
                data_source_id,
                manifest["version"],
                checksum,
                len(points),
                added_count,
                updated_count,
                missing_count,
                details,
            ),
        )
        db.execute(
            """
            INSERT INTO audit_logs (username, action, details)
            VALUES ('system', 'sync_device_points', ?)
            """,
            (f"Point catalog {manifest['version']}: {details}",),
        )
        db.commit()
        return {
            "status": "updated",
            "version": manifest["version"],
            "record_count": len(points),
            "added_count": added_count,
            "updated_count": updated_count,
            "missing_count": missing_count,
        }
    except Exception as exc:
        db.rollback()
        app.logger.exception("Cannot synchronize bundled point catalog: %s", exc)
        try:
            db.execute(
                """
                INSERT INTO point_catalog_versions (
                    data_source_id, version, checksum, record_count, status, details
                ) VALUES (?, ?, ?, ?, 'failed', ?)
                """,
                (data_source_id, manifest.get("version", "unknown"), checksum, len(points), str(exc)[:1000]),
            )
            db.commit()
        except sqlite3.Error:
            db.rollback()
        return {"status": "failed", "details": str(exc)}


@app.before_request
def before_request() -> None:
    global _INITIALIZED_DB_PATH
    resolved_db_path = DB_PATH.resolve()
    if _INITIALIZED_DB_PATH == resolved_db_path:
        return
    with _DB_INIT_LOCK:
        if _INITIALIZED_DB_PATH != resolved_db_path:
            init_db()
            _INITIALIZED_DB_PATH = resolved_db_path


def current_user() -> sqlite3.Row | None:
    user_id = session.get("user_id")
    if not user_id:
        return None
    return get_db().execute(
        "SELECT * FROM users WHERE id = ? AND active = 1",
        (user_id,),
    ).fetchone()


def login_required():
    if current_user() is None:
        flash("請先登入。", "warning")
        return redirect(url_for("login"))
    return None


def admin_required():
    guard = login_required()
    if guard:
        return guard
    if session.get("role") != "admin":
        flash("此功能限管理者使用。", "danger")
        return redirect(url_for("fill_report"))
    return None


def report_access_required():
    guard = login_required()
    if guard:
        return guard
    if session.get("role") not in {"admin", "viewer"}:
        flash("此功能限管理者或訪客使用。", "danger")
        return redirect(url_for("fill_report"))
    return None


def active_energy_types() -> list[sqlite3.Row]:
    return get_db().execute(
        """
        SELECT * FROM energy_types
        WHERE active = 1
        ORDER BY display_order, id
        """
    ).fetchall()


def active_areas() -> list[sqlite3.Row]:
    return get_db().execute(
        """
        SELECT * FROM areas
        WHERE active = 1
        ORDER BY display_order, id
        """
    ).fetchall()


@app.route("/", methods=["GET"])
def index():
    if not session.get("user_id"):
        return redirect(url_for("login"))
    if session.get("role") == "admin":
        return redirect(url_for("admin_dashboard"))
    if session.get("role") == "viewer":
        return redirect(url_for("admin_reports"))
    return redirect(url_for("fill_report"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        user = get_db().execute(
            """
            SELECT * FROM users
            WHERE username = ? AND password = ? AND active = 1
            """,
            (username, password),
        ).fetchone()

        if user:
            language = session.get("language", "zh-TW")
            session.clear()
            session["language"] = language
            session["user_id"] = user["id"]
            session["username"] = user["username"]
            session["display_name"] = user["display_name"]
            session["role"] = "viewer" if user["viewer"] else user["role"]
            flash("登入成功。", "success")
            log_action("login", f"Successful login from role: {session['role']}")
            return redirect(url_for("index"))

        flash("帳號或密碼錯誤。", "danger")

    return render_template("login.html")


@app.route("/logout", methods=["POST"])
def logout():
    log_action("logout")
    language = session.get("language", "zh-TW")
    session.clear()
    session["language"] = language
    flash("已登出。", "info")
    return redirect(url_for("login"))


@app.route("/admin", methods=["GET"])
def admin_dashboard():
    guard = admin_required()
    if guard:
        return guard

    db = get_db()
    report_month = request.args.get("report_month", "").strip() or latest_report_month()
    if not valid_report_month(report_month):
        flash("分析月份格式不正確。", "danger")
        return redirect(url_for("admin_dashboard", report_month=latest_report_month()))

    previous_month = previous_report_month(report_month)
    current_amount = db.execute(
        "SELECT COALESCE(SUM(amount), 0) FROM energy_entries WHERE report_month = ?",
        (report_month,),
    ).fetchone()[0]
    previous_amount = db.execute(
        "SELECT COALESCE(SUM(amount), 0) FROM energy_entries WHERE report_month = ?",
        (previous_month,),
    ).fetchone()[0]
    month_change = (
        ((current_amount - previous_amount) / previous_amount) * 100
        if previous_amount
        else None
    )
    active_fillers = db.execute(
        "SELECT COUNT(*) FROM users WHERE role = 'filler' AND viewer = 0 AND active = 1"
    ).fetchone()[0]
    submitted_fillers = db.execute(
        """
        SELECT COUNT(DISTINCT rs.user_id)
        FROM report_submissions rs
        JOIN users u ON u.id = rs.user_id
        WHERE rs.report_month = ? AND u.role = 'filler' AND u.viewer = 0 AND u.active = 1
        """,
        (report_month,),
    ).fetchone()[0]
    valid_entries = db.execute(
        """
        SELECT COUNT(*) FROM energy_entries
        WHERE report_month = ? AND (quantity IS NOT NULL OR amount IS NOT NULL)
        """,
        (report_month,),
    ).fetchone()[0]

    monthly_rows = db.execute(
        """
        SELECT report_month, COALESCE(SUM(amount), 0) AS total_amount
        FROM energy_entries
        GROUP BY report_month
        ORDER BY report_month DESC
        LIMIT 12
        """
    ).fetchall()[::-1]
    max_month_amount = max((row["total_amount"] for row in monthly_rows), default=0)
    monthly_trend = [
        {
            "month": row["report_month"],
            "amount": row["total_amount"],
            "percent": (row["total_amount"] / max_month_amount * 100) if max_month_amount else 0,
            "selected": row["report_month"] == report_month,
        }
        for row in monthly_rows
    ]

    energy_rows = db.execute(
        """
        SELECT et.name, COALESCE(SUM(ee.amount), 0) AS total_amount
        FROM energy_entries ee
        JOIN energy_types et ON et.id = ee.energy_type_id
        WHERE ee.report_month = ? AND ee.amount IS NOT NULL
        GROUP BY et.id, et.name
        HAVING SUM(ee.amount) != 0
        ORDER BY total_amount DESC
        """,
        (report_month,),
    ).fetchall()
    energy_total = sum(row["total_amount"] for row in energy_rows)
    energy_breakdown = [
        {
            "name": row["name"],
            "amount": row["total_amount"],
            "percent": (row["total_amount"] / energy_total * 100) if energy_total else 0,
        }
        for row in energy_rows
    ]

    area_rows = db.execute(
        """
        SELECT a.name, COALESCE(SUM(ee.amount), 0) AS total_amount
        FROM areas a
        LEFT JOIN energy_entries ee
            ON ee.area_id = a.id AND ee.report_month = ? AND ee.amount IS NOT NULL
        WHERE a.active = 1
        GROUP BY a.id, a.name
        ORDER BY total_amount DESC, a.display_order, a.id
        """,
        (report_month,),
    ).fetchall()
    max_area_amount = max((row["total_amount"] for row in area_rows), default=0)
    area_ranking = [
        {
            "name": row["name"],
            "amount": row["total_amount"],
            "percent": (row["total_amount"] / max_area_amount * 100) if max_area_amount else 0,
            "share": (row["total_amount"] / current_amount * 100) if current_amount else 0,
        }
        for row in area_rows
    ]
    stats = {
        "fillers": db.execute("SELECT COUNT(*) FROM users WHERE role = 'filler' AND viewer = 0").fetchone()[0],
        "energy_types": db.execute("SELECT COUNT(*) FROM energy_types WHERE active = 1").fetchone()[0],
        "areas": db.execute("SELECT COUNT(*) FROM areas WHERE active = 1").fetchone()[0],
        "entries": db.execute("SELECT COUNT(*) FROM energy_entries").fetchone()[0],
    }
    return render_template(
        "admin_dashboard.html",
        stats=stats,
        report_month=report_month,
        previous_month=previous_month,
        current_amount=current_amount,
        previous_amount=previous_amount,
        month_change=month_change,
        active_fillers=active_fillers,
        submitted_fillers=submitted_fillers,
        pending_fillers=max(active_fillers - submitted_fillers, 0),
        valid_entries=valid_entries,
        monthly_trend=monthly_trend,
        energy_breakdown=energy_breakdown,
        area_ranking=area_ranking,
        audit_logs=db.execute("SELECT * FROM audit_logs ORDER BY id DESC LIMIT 15").fetchall(),
    )


@app.route("/admin/backup", methods=["GET"])
def admin_backup():
    guard = admin_required()
    if guard:
        return guard
    try:
        log_action("backup_db", "Downloaded database backup")
        return send_file(
            DB_PATH,
            as_attachment=True,
            download_name=f"energy_system_backup_{date.today().strftime('%Y%m%d')}.db",
            mimetype="application/x-sqlite3"
        )
    except Exception as e:
        flash(f"備份失敗: {str(e)}", "danger")
        return redirect(url_for("admin_dashboard"))


@app.route("/admin/users", methods=["GET", "POST"])
def admin_users():
    guard = admin_required()
    if guard:
        return guard

    db = get_db()
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "").strip()
        display_name = request.form.get("display_name", "").strip() or username
        account_type = request.form.get("account_type", "filler")
        if account_type not in {"filler", "viewer"}:
            account_type = "filler"
        if not username or not password:
            flash("請輸入帳號與密碼。", "danger")
        else:
            try:
                db.execute(
                    """
                    INSERT INTO users (username, password, display_name, role, viewer)
                    VALUES (?, ?, ?, 'filler', ?)
                    """,
                    (username, password, display_name, account_type == "viewer"),
                )
                user_id = db.execute(
                    "SELECT id FROM users WHERE username = ?",
                    (username,),
                ).fetchone()["id"]
                if account_type == "filler":
                    db.execute(
                        """
                        INSERT INTO user_area_permissions (user_id, area_id)
                        SELECT ?, id FROM areas WHERE active = 1
                        """,
                        (user_id,),
                    )
                db.commit()
                log_action("create_user", f"Created account {username} as {account_type}")
                flash("填寫員已建立。", "success")
            except sqlite3.IntegrityError:
                flash("帳號已存在。", "danger")
        return redirect(url_for("admin_users"))

    users = db.execute(
        """
        SELECT * FROM users
        WHERE role = 'filler'
        ORDER BY id DESC
        """
    ).fetchall()
    energy_permissions = {
        row["user_id"]: set(row["energy_ids"].split(",") if row["energy_ids"] else [])
        for row in db.execute(
            """
            SELECT user_id, GROUP_CONCAT(energy_type_id) AS energy_ids
            FROM user_energy_permissions
            GROUP BY user_id
            """
        ).fetchall()
    }
    area_permissions = {
        row["user_id"]: set(row["area_ids"].split(",") if row["area_ids"] else [])
        for row in db.execute(
            """
            SELECT user_id, GROUP_CONCAT(area_id) AS area_ids
            FROM user_area_permissions
            GROUP BY user_id
            """
        ).fetchall()
    }
    return render_template(
        "admin_users.html",
        users=users,
        energy_types=active_energy_types(),
        areas=active_areas(),
        energy_permissions=energy_permissions,
        area_permissions=area_permissions,
    )


@app.route("/admin/users/<int:user_id>/permissions", methods=["POST"])
def update_user_permissions(user_id: int):
    guard = admin_required()
    if guard:
        return guard

    db = get_db()
    user = db.execute(
        "SELECT * FROM users WHERE id = ? AND role = 'filler' AND viewer = 0",
        (user_id,),
    ).fetchone()
    if user is None:
        flash("找不到填寫員。", "warning")
        return redirect(url_for("admin_users"))

    selected_energy_ids = request.form.getlist("energy_type_ids")
    selected_area_ids = request.form.getlist("area_ids")
    db.execute("DELETE FROM user_energy_permissions WHERE user_id = ?", (user_id,))
    db.executemany(
        """
        INSERT INTO user_energy_permissions (user_id, energy_type_id)
        VALUES (?, ?)
        """,
        [(user_id, energy_id) for energy_id in selected_energy_ids],
    )
    db.execute("DELETE FROM user_area_permissions WHERE user_id = ?", (user_id,))
    db.executemany(
        """
        INSERT INTO user_area_permissions (user_id, area_id)
        VALUES (?, ?)
        """,
        [(user_id, area_id) for area_id in selected_area_ids],
    )
    db.commit()
    log_action("update_permissions", f"Updated Y/X permissions for filler: {user['username']}")
    flash("填寫權限已更新。", "success")
    return redirect(url_for("admin_users"))


@app.route("/admin/users/<int:user_id>/password", methods=["POST"])
def reset_user_password(user_id: int):
    guard = admin_required()
    if guard:
        return guard

    new_password = request.form.get("new_password", "").strip()
    if not new_password:
        flash("請輸入新密碼。", "danger")
        return redirect(url_for("admin_users"))

    result = get_db().execute(
        """
        UPDATE users
        SET password = ?
        WHERE id = ? AND role = 'filler'
        """,
        (new_password, user_id),
    )
    get_db().commit()
    if result.rowcount:
        log_action("reset_password", f"Reset password for filler ID: {user_id}")
        flash("密碼已重設。", "success")
    else:
        flash("找不到填寫員。", "warning")
    return redirect(url_for("admin_users"))


@app.route("/admin/settings", methods=["GET", "POST"])
def admin_settings():
    guard = admin_required()
    if guard:
        return guard

    db = get_db()
    action = request.form.get("action")
    if request.method == "POST" and action == "add_energy":
        name = request.form.get("name", "").strip()
        unit = request.form.get("unit", "").strip()
        if not name or not unit:
            flash("請輸入能源種類與單位。", "danger")
        else:
            try:
                next_order = db.execute("SELECT COALESCE(MAX(display_order), 0) + 1 FROM energy_types").fetchone()[0]
                db.execute(
                    "INSERT INTO energy_types (name, unit, display_order) VALUES (?, ?, ?)",
                    (name, unit, next_order),
                )
                db.commit()
                log_action("add_energy_type", f"Added: {name} (Unit: {unit})")
                flash("能源種類已新增。", "success")
            except sqlite3.IntegrityError:
                flash("能源種類已存在。", "danger")
        return redirect(url_for("admin_settings"))

    if request.method == "POST" and action == "add_area":
        name = request.form.get("name", "").strip()
        if not name:
            flash("請輸入區域名稱。", "danger")
        else:
            try:
                next_order = db.execute("SELECT COALESCE(MAX(display_order), 0) + 1 FROM areas").fetchone()[0]
                db.execute(
                    "INSERT INTO areas (name, display_order) VALUES (?, ?)",
                    (name, next_order),
                )
                db.commit()
                log_action("add_area", f"Added area: {name}")
                flash("區域已新增。", "success")
            except sqlite3.IntegrityError:
                flash("區域已存在。", "danger")
        return redirect(url_for("admin_settings"))

    return render_template(
        "admin_settings.html",
        energy_types=active_energy_types(),
        areas=active_areas(),
    )


@app.route("/admin/control-limits", methods=["GET", "POST"])
def admin_control_limits():
    guard = admin_required()
    if guard:
        return guard

    db = get_db()
    if request.method == "POST" and request.form.get("action") == "recalculate":
        pct_val = request.form.get("percentage", "10")
        try:
            percentage = float(pct_val) / 100.0
        except ValueError:
            percentage = 0.1
        seed_control_limits(db, replace=True, percentage=percentage)
        db.commit()
        log_action("recalculate_control_limits", f"Recalculated with deviation ±{pct_val}%")
        flash("管制值已依歷史平均值重新計算。", "success")
        return redirect(url_for("admin_control_limits"))

    if request.method == "POST":
        rows = db.execute("SELECT id FROM control_limits ORDER BY id").fetchall()
        updates = []
        try:
            for row in rows:
                control_id = row["id"]
                quantity_enabled = f"quantity_enabled_{control_id}" in request.form
                amount_enabled = f"amount_enabled_{control_id}" in request.form

                def optional_float(field: str) -> float | None:
                    value = request.form.get(f"{field}_{control_id}", "").strip()
                    return float(value) if value else None

                quantity_min = optional_float("quantity_min")
                quantity_max = optional_float("quantity_max")
                amount_min = optional_float("amount_min")
                amount_max = optional_float("amount_max")
                if quantity_enabled and (quantity_min is None or quantity_max is None):
                    raise ValueError
                if amount_enabled and (amount_min is None or amount_max is None):
                    raise ValueError
                if quantity_min is not None and quantity_max is not None and quantity_min > quantity_max:
                    flash("下限不可大於上限。", "danger")
                    return redirect(url_for("admin_control_limits"))
                if amount_min is not None and amount_max is not None and amount_min > amount_max:
                    flash("下限不可大於上限。", "danger")
                    return redirect(url_for("admin_control_limits"))
                updates.append((
                    quantity_enabled, quantity_min, quantity_max,
                    amount_enabled, amount_min, amount_max, control_id,
                ))
        except ValueError:
            flash("管制值格式不正確，請確認上下限。", "danger")
            return redirect(url_for("admin_control_limits"))

        db.executemany(
            """
            UPDATE control_limits
            SET quantity_enabled = ?, quantity_min = ?, quantity_max = ?,
                amount_enabled = ?, amount_min = ?, amount_max = ?,
                updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
            """,
            updates,
        )
        db.commit()
        log_action("update_control_limits", "Manually updated control limit grid")
        flash("管制值已更新。", "success")
        return redirect(url_for("admin_control_limits"))

    rows = db.execute(
        """
        SELECT cl.*, et.name AS energy_name, et.unit, a.name AS area_name
        FROM control_limits cl
        JOIN energy_types et ON et.id = cl.energy_type_id
        JOIN areas a ON a.id = cl.area_id
        WHERE et.active = 1 AND a.active = 1
        ORDER BY et.display_order, et.id, a.display_order, a.id
        """
    ).fetchall()
    return render_template("admin_control_limits.html", rows=rows)


EMS_VALUE_TYPES = {
    "instantaneous": "瞬時值",
    "cumulative": "累積表值",
    "state": "狀態",
    "calculated": "計算值",
}
EMS_AGGREGATION_METHODS = {
    "average": "平均",
    "delta": "期末減期初",
    "sum": "加總",
    "max": "最大值",
    "min": "最小值",
    "last": "最後值",
    "duration": "運轉時間",
    "none": "不統計",
}
EMS_FLOW_DIRECTIONS = ["輸入", "消耗", "生產", "輸出", "回收", "排放"]
EMS_ALARM_SEVERITIES = {"info": "提示", "warning": "警告", "critical": "嚴重"}
EMS_ALARM_VALUE_MODES = {
    "raw": "原始讀值",
    "interval_delta": "每次收集差值",
    "change_rate": "變化率",
}


def ems_lookup_rows(table: str) -> list[sqlite3.Row]:
    allowed_tables = {
        "ems_energy_categories",
        "ems_measurement_types",
        "ems_units",
        "ems_locations",
    }
    if table not in allowed_tables:
        raise ValueError("Invalid EMS lookup table")
    return get_db().execute(
        f"SELECT id, name FROM {table} WHERE active = 1 ORDER BY display_order, id"
    ).fetchall()


def ems_complete_sql(alias: str = "cfg") -> str:
    return f"""(
        {alias}.energy_category_id IS NOT NULL
        AND {alias}.measurement_type_id IS NOT NULL
        AND {alias}.unit_id IS NOT NULL
        AND {alias}.location_id IS NOT NULL
        AND COALESCE({alias}.value_type, '') <> ''
        AND COALESCE({alias}.aggregation_method, '') <> ''
    )"""


def ems_issue_sql(point_alias: str = "dp") -> str:
    return f"""(
        {point_alias}.source_present = 0
        OR TRIM(COALESCE({point_alias}.point_name, '')) = ''
        OR TRIM(COALESCE({point_alias}.json_path, '')) = ''
        OR (
            TRIM(COALESCE({point_alias}.point_name, '')) <> ''
            AND (SELECT COUNT(*) FROM device_data_points duplicate
                 WHERE duplicate.data_source_id = {point_alias}.data_source_id
                   AND duplicate.point_name = {point_alias}.point_name) > 1
        )
    )"""


def ems_point_filters() -> tuple[list[str], list, dict]:
    query = request.args.get("q", "").strip()
    monitor = request.args.get("monitor", "all")
    category_id = request.args.get("category_id", "").strip()
    location_id = request.args.get("location_id", "").strip()
    status = request.args.get("status", "all")
    filters = ["dp.data_source_id = ?"]
    params: list = [1]
    if query:
        filters.append(
            "(dp.point_name LIKE ? OR dp.json_path LIKE ? OR dp.description LIKE ? OR cfg.display_name LIKE ?)"
        )
        keyword = f"%{query}%"
        params.extend([keyword, keyword, keyword, keyword])
    if monitor == "on":
        filters.append("cfg.monitor_enabled = 1")
    elif monitor == "off":
        filters.append("cfg.monitor_enabled = 0")
    if category_id.isdigit():
        filters.append("cfg.energy_category_id = ?")
        params.append(int(category_id))
    if location_id.isdigit():
        filters.append("cfg.location_id = ?")
        params.append(int(location_id))
    if status == "complete":
        filters.append(ems_complete_sql())
    elif status == "incomplete":
        filters.append(f"NOT {ems_complete_sql()}")
    elif status == "issue":
        filters.append(ems_issue_sql())
    elif status == "no_limit":
        filters.append("cfg.lower_limit IS NULL AND cfg.upper_limit IS NULL")
    elif status == "source_missing":
        filters.append("dp.source_present = 0")
    values = {
        "q": query,
        "monitor": monitor,
        "category_id": category_id,
        "location_id": location_id,
        "status": status,
    }
    return filters, params, values


def ems_point_select_sql(where_sql: str) -> str:
    return f"""
        SELECT
            dp.id, dp.external_id, dp.data_source_id, dp.point_name, dp.json_path,
            dp.description, dp.source_is_active, dp.point_type,
            dp.fetch_interval_minutes, dp.source_present, dp.last_synced_at,
            ds.name AS source_name,
            cfg.monitor_enabled, cfg.lower_limit, cfg.upper_limit, cfg.display_name,
            cfg.energy_category_id, cfg.measurement_type_id, cfg.unit_id,
            cfg.location_id, cfg.equipment_name, cfg.value_type,
            cfg.aggregation_method, cfg.flow_direction, cfg.alarm_severity,
            cfg.alarm_value_mode, cfg.alarm_consecutive_samples,
            cfg.alarm_deadband, cfg.notes,
            cfg.updated_at, updater.username AS updated_by_username,
            category.name AS category_name,
            measurement.name AS measurement_name,
            unit.name AS unit_name,
            location.name AS location_name,
            (SELECT COUNT(*) FROM device_data_points duplicate
             WHERE duplicate.data_source_id = dp.data_source_id
               AND duplicate.point_name = dp.point_name
               AND TRIM(COALESCE(dp.point_name, '')) <> '') AS duplicate_count
        FROM device_data_points dp
        JOIN device_sources ds ON ds.external_id = dp.data_source_id
        JOIN ems_point_configs cfg ON cfg.point_id = dp.id
        LEFT JOIN ems_energy_categories category ON category.id = cfg.energy_category_id
        LEFT JOIN ems_measurement_types measurement ON measurement.id = cfg.measurement_type_id
        LEFT JOIN ems_units unit ON unit.id = cfg.unit_id
        LEFT JOIN ems_locations location ON location.id = cfg.location_id
        LEFT JOIN users updater ON updater.id = cfg.updated_by
        WHERE {where_sql}
    """


@app.route("/admin/ems-points", methods=["GET"])
def admin_ems_points():
    guard = admin_required()
    if guard:
        return guard

    db = get_db()
    filters, params, filter_values = ems_point_filters()
    active_filters = []
    if filter_values["q"]:
        active_filters.append(("搜尋點位", filter_values["q"]))
    if filter_values["monitor"] in {"on", "off"}:
        active_filters.append(("監控狀態", translate("監控中") if filter_values["monitor"] == "on" else translate("不監控")))
    status_labels = {
        "complete": "分類完成", "incomplete": "尚未分類", "issue": "來源有問題",
        "no_limit": "尚未設定門檻", "source_missing": "來源已不存在",
    }
    if filter_values["status"] in status_labels:
        active_filters.append(("資料狀態", translate(status_labels[filter_values["status"]])))
    for key, table, label in [("category_id", "ems_energy_categories", "能源類別"), ("location_id", "ems_locations", "區域")]:
        if filter_values[key].isdigit():
            lookup = db.execute(f"SELECT name FROM {table} WHERE id = ?", (int(filter_values[key]),)).fetchone()
            value = lookup[0] if lookup else filter_values[key]
            active_filters.append((label, translate(value) if key == "category_id" else value))
    where_sql = " AND ".join(filters)
    try:
        page = max(int(request.args.get("page", "1")), 1)
    except ValueError:
        page = 1
    try:
        per_page = int(request.args.get("per_page", "50"))
    except ValueError:
        per_page = 50
    if per_page not in {25, 50, 100}:
        per_page = 50
    total_rows = db.execute(
        f"SELECT COUNT(*) FROM ({ems_point_select_sql(where_sql)}) filtered",
        params,
    ).fetchone()[0]
    total_pages = max(math.ceil(total_rows / per_page), 1)
    page = min(page, total_pages)
    point_rows = db.execute(
        ems_point_select_sql(where_sql)
        + " ORDER BY dp.source_present DESC, dp.external_id LIMIT ? OFFSET ?",
        [*params, per_page, (page - 1) * per_page],
    ).fetchall()

    complete_expression = ems_complete_sql()
    issue_expression = ems_issue_sql()
    stats = db.execute(
        f"""
        SELECT
            COUNT(*) AS total,
            SUM(CASE WHEN cfg.monitor_enabled = 1 THEN 1 ELSE 0 END) AS monitored,
            SUM(CASE WHEN cfg.lower_limit IS NOT NULL OR cfg.upper_limit IS NOT NULL THEN 1 ELSE 0 END) AS with_limits,
            SUM(CASE WHEN NOT {complete_expression} THEN 1 ELSE 0 END) AS incomplete,
            SUM(CASE WHEN {issue_expression} THEN 1 ELSE 0 END) AS issues
        FROM device_data_points dp
        JOIN ems_point_configs cfg ON cfg.point_id = dp.id
        WHERE dp.data_source_id = 1
        """
    ).fetchone()
    catalog_version = db.execute(
        """
        SELECT * FROM point_catalog_versions
        WHERE data_source_id = 1
        ORDER BY id DESC LIMIT 1
        """
    ).fetchone()
    return render_template(
        "admin_ems_points.html",
        points=point_rows,
        point_payloads=[dict(row) for row in point_rows],
        stats=stats,
        filters=filter_values,
        active_filters=active_filters,
        categories=ems_lookup_rows("ems_energy_categories"),
        measurement_types=ems_lookup_rows("ems_measurement_types"),
        units=ems_lookup_rows("ems_units"),
        locations=ems_lookup_rows("ems_locations"),
        value_types=EMS_VALUE_TYPES,
        aggregation_methods=EMS_AGGREGATION_METHODS,
        flow_directions=EMS_FLOW_DIRECTIONS,
        alarm_severities=EMS_ALARM_SEVERITIES,
        alarm_value_modes=EMS_ALARM_VALUE_MODES,
        catalog_version=catalog_version,
        page=page,
        per_page=per_page,
        total_rows=total_rows,
        total_pages=total_pages,
    )


def optional_finite_float(field: str, label: str) -> float | None:
    raw = request.form.get(field, "").strip()
    if not raw:
        return None
    try:
        value = float(raw)
    except ValueError as exc:
        raise ValueError(f"{label}必須是有效數字。") from exc
    if not math.isfinite(value):
        raise ValueError(f"{label}必須是有限數字。")
    return value


def optional_lookup_id(db: sqlite3.Connection, field: str, table: str, label: str) -> int | None:
    raw = request.form.get(field, "").strip()
    if not raw:
        return None
    if not raw.isdigit():
        raise ValueError(f"{label}選項不正確。")
    lookup_id = int(raw)
    if db.execute(f"SELECT 1 FROM {table} WHERE id = ? AND active = 1", (lookup_id,)).fetchone() is None:
        raise ValueError(f"{label}選項不存在。")
    return lookup_id


def safe_return_url(default_endpoint: str = "admin_ems_points") -> str:
    return_url = request.form.get("return_to", "").strip()
    if not return_url.startswith("/") or return_url.startswith("//"):
        return url_for(default_endpoint)
    return return_url


@app.route("/admin/ems-points/<int:point_id>", methods=["POST"])
def update_ems_point(point_id: int):
    guard = admin_required()
    if guard:
        return guard

    db = get_db()
    point = db.execute(
        """
        SELECT dp.*, cfg.id AS config_id
        FROM device_data_points dp
        JOIN ems_point_configs cfg ON cfg.point_id = dp.id
        WHERE dp.id = ? AND dp.data_source_id = 1
        """,
        (point_id,),
    ).fetchone()
    if point is None:
        flash("找不到指定的 EMS 點位。", "danger")
        return redirect(safe_return_url())

    try:
        lower_limit = optional_finite_float("lower_limit", "下限")
        upper_limit = optional_finite_float("upper_limit", "上限")
        deadband = optional_finite_float("alarm_deadband", "告警緩衝值")
        if lower_limit is not None and upper_limit is not None and lower_limit > upper_limit:
            raise ValueError("下限不可大於上限。")
        if deadband is not None and deadband < 0:
            raise ValueError("告警緩衝值不可小於零。")
        consecutive_raw = request.form.get("alarm_consecutive_samples", "1").strip()
        try:
            consecutive_samples = int(consecutive_raw or "1")
        except (ValueError, OverflowError) as exc:
            raise ValueError("連續異常次數必須介於 1 到 168 次。") from exc
        if consecutive_samples < 1 or consecutive_samples > 168:
            raise ValueError("連續異常次數必須介於 1 到 168 次。")
        category_id = optional_lookup_id(db, "energy_category_id", "ems_energy_categories", "能源類別")
        measurement_type_id = optional_lookup_id(db, "measurement_type_id", "ems_measurement_types", "測量類型")
        unit_id = optional_lookup_id(db, "unit_id", "ems_units", "工程單位")
        location_id = optional_lookup_id(db, "location_id", "ems_locations", "區域")
        value_type = request.form.get("value_type", "").strip() or None
        aggregation_method = request.form.get("aggregation_method", "").strip() or None
        flow_direction = request.form.get("flow_direction", "").strip() or None
        severity = request.form.get("alarm_severity", "warning").strip()
        alarm_value_mode = request.form.get("alarm_value_mode", "raw").strip()
        if value_type is not None and value_type not in EMS_VALUE_TYPES:
            raise ValueError("數值性質選項不正確。")
        if aggregation_method is not None and aggregation_method not in EMS_AGGREGATION_METHODS:
            raise ValueError("統計方式選項不正確。")
        if flow_direction is not None and flow_direction not in EMS_FLOW_DIRECTIONS:
            raise ValueError("流向選項不正確。")
        if severity not in EMS_ALARM_SEVERITIES:
            raise ValueError("告警等級選項不正確。")
        if alarm_value_mode not in EMS_ALARM_VALUE_MODES:
            raise ValueError("告警判斷值選項不正確。")
    except (ValueError, OverflowError) as exc:
        flash(str(exc), "danger")
        return redirect(safe_return_url())

    db.execute(
        """
        UPDATE ems_point_configs
        SET monitor_enabled = ?, lower_limit = ?, upper_limit = ?, display_name = ?,
            energy_category_id = ?, measurement_type_id = ?, unit_id = ?, location_id = ?,
            equipment_name = ?, value_type = ?, aggregation_method = ?, flow_direction = ?,
            alarm_severity = ?, alarm_value_mode = ?, alarm_consecutive_samples = ?,
            alarm_deadband = ?, notes = ?,
            updated_by = ?, updated_at = CURRENT_TIMESTAMP
        WHERE point_id = ?
        """,
        (
            "monitor_enabled" in request.form,
            lower_limit,
            upper_limit,
            request.form.get("display_name", "").strip() or None,
            category_id,
            measurement_type_id,
            unit_id,
            location_id,
            request.form.get("equipment_name", "").strip() or None,
            value_type,
            aggregation_method,
            flow_direction,
            severity,
            alarm_value_mode,
            consecutive_samples,
            deadband or 0,
            request.form.get("notes", "").strip() or None,
            session["user_id"],
            point_id,
        ),
    )
    db.commit()
    log_action(
        "update_ems_point",
        f"Updated EMS point {point['external_id']} ({point['point_name'] or 'blank PointName'})",
    )
    flash(f"點位 {point['external_id']} 的 EMS 設定已儲存。", "success")
    return redirect(safe_return_url())


@app.route("/admin/ems-points/bulk", methods=["POST"])
def bulk_update_ems_points():
    guard = admin_required()
    if guard:
        return guard

    selected_ids = []
    for raw_id in request.form.getlist("point_ids")[:500]:
        if raw_id.isdigit():
            selected_ids.append(int(raw_id))
    selected_ids = sorted(set(selected_ids))
    if not selected_ids:
        flash("請先選擇至少一個點位。", "warning")
        return redirect(safe_return_url())

    db = get_db()
    placeholders = ",".join("?" for _ in selected_ids)
    allowed_ids = [
        row["id"] for row in db.execute(
            f"SELECT id FROM device_data_points WHERE data_source_id = 1 AND id IN ({placeholders})",
            selected_ids,
        ).fetchall()
    ]
    if not allowed_ids:
        flash("選取的點位不存在。", "danger")
        return redirect(safe_return_url())
    placeholders = ",".join("?" for _ in allowed_ids)
    action = request.form.get("action", "")
    if action in {"monitor_on", "monitor_off"}:
        monitor_enabled = action == "monitor_on"
        db.execute(
            f"""
            UPDATE ems_point_configs
            SET monitor_enabled = ?, updated_by = ?, updated_at = CURRENT_TIMESTAMP
            WHERE point_id IN ({placeholders})
            """,
            [monitor_enabled, session["user_id"], *allowed_ids],
        )
        action_label = "啟用監控" if monitor_enabled else "停止監控"
    elif action == "set_category":
        try:
            category_id = optional_lookup_id(db, "bulk_category_id", "ems_energy_categories", "能源類別")
        except ValueError as exc:
            flash(str(exc), "danger")
            return redirect(safe_return_url())
        if category_id is None:
            flash("請選擇要套用的能源類別。", "warning")
            return redirect(safe_return_url())
        db.execute(
            f"""
            UPDATE ems_point_configs
            SET energy_category_id = ?, updated_by = ?, updated_at = CURRENT_TIMESTAMP
            WHERE point_id IN ({placeholders})
            """,
            [category_id, session["user_id"], *allowed_ids],
        )
        action_label = "設定能源類別"
    else:
        flash("批次操作不正確。", "danger")
        return redirect(safe_return_url())
    db.commit()
    log_action("bulk_update_ems_points", f"{action_label}: {len(allowed_ids)} points")
    flash(f"已為 {len(allowed_ids)} 個點位完成「{action_label}」。", "success")
    return redirect(safe_return_url())


@app.route("/admin/ems-points/export", methods=["GET"])
def export_ems_points():
    guard = admin_required()
    if guard:
        return guard

    filters, params, _ = ems_point_filters()
    rows = get_db().execute(
        ems_point_select_sql(" AND ".join(filters)) + " ORDER BY dp.external_id",
        params,
    ).fetchall()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "EMS Point Config"
    sheet.append([translate(label) for label in [
        "DataSourceId", "外部 Id", "PointName", "JsonPath", "來源說明",
        "來源存在", "是否監控", "下限", "上限", "顯示名稱",
        "能源類別", "測量類型", "工程單位", "區域", "設備／系統",
        "數值性質", "統計方式", "流向", "告警等級",
        "告警判斷值", "連續異常次數", "告警緩衝值", "備註",
        "更新者", "更新時間",
    ]])
    for row in rows:
        sheet.append([
            row["data_source_id"], row["external_id"], row["point_name"], row["json_path"],
            row["description"], bool(row["source_present"]), bool(row["monitor_enabled"]),
            row["lower_limit"], row["upper_limit"], row["display_name"], row["category_name"],
            row["measurement_name"], row["unit_name"], row["location_name"],
            row["equipment_name"], row["value_type"], row["aggregation_method"],
            row["flow_direction"], row["alarm_severity"], row["alarm_value_mode"],
            row["alarm_consecutive_samples"], row["alarm_deadband"], row["notes"],
            row["updated_by_username"], row["updated_at"],
        ])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(
            max(len(str(cell.value or "")) for cell in column) + 2,
            42,
        )
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name=f"ems-point-config-{date.today().isoformat()}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/fill", methods=["GET", "POST"])
def fill_report():
    guard = login_required()
    if guard:
        return guard
    if session.get("role") == "admin":
        return redirect(url_for("admin_dashboard"))
    if session.get("role") == "viewer":
        return redirect(url_for("admin_reports"))

    db = get_db()
    user_id = session["user_id"]
    report_month = (
        request.form.get("report_month", "").strip()
        if request.method == "POST"
        else request.args.get("report_month", "").strip()
    ) or latest_report_month()
    if not valid_report_month(report_month):
        flash("填報月份格式不正確。", "danger")
        return redirect(url_for("fill_report", report_month=latest_report_month()))
    max_report_month = latest_report_month()
    if report_month > max_report_month:
        flash("填報月份只能選擇上個月或更早月份。", "danger")
        return redirect(url_for("fill_report", report_month=max_report_month))

    energy_types = db.execute(
        """
        SELECT et.*
        FROM energy_types et
        JOIN user_energy_permissions p ON p.energy_type_id = et.id
        WHERE p.user_id = ? AND et.active = 1
        ORDER BY et.display_order, et.id
        """,
        (user_id,),
    ).fetchall()
    areas = db.execute(
        """
        SELECT a.*
        FROM areas a
        JOIN user_area_permissions p ON p.area_id = a.id
        WHERE p.user_id = ? AND a.active = 1
        ORDER BY a.display_order, a.id
        """,
        (user_id,),
    ).fetchall()
    control_limits = {
        (row["energy_type_id"], row["area_id"]): row
        for row in db.execute(
            """
            SELECT cl.*
            FROM control_limits cl
            JOIN user_energy_permissions ep
                ON ep.energy_type_id = cl.energy_type_id AND ep.user_id = ?
            JOIN user_area_permissions ap
                ON ap.area_id = cl.area_id AND ap.user_id = ?
            """,
            (user_id, user_id),
        ).fetchall()
    }
    submission = db.execute(
        """
        SELECT * FROM report_submissions
        WHERE report_month = ? AND user_id = ?
        """,
        (report_month, user_id),
    ).fetchone()
    is_submitted = submission is not None

    if request.method == "POST":
        if not energy_types or not areas:
            flash("目前尚未被指派可填寫的能源種類或單位，無法送出。", "warning")
            return redirect(url_for("fill_report", report_month=report_month))

        control_warning_count = 0
        for energy_type in energy_types:
            for area in areas:
                quantity = request.form.get(f"quantity_{energy_type['id']}_{area['id']}", "").strip()
                amount = request.form.get(f"amount_{energy_type['id']}_{area['id']}", "").strip()
                limit = control_limits.get((energy_type["id"], area["id"]))
                for value, enabled_key, min_key, max_key in (
                    (quantity, "quantity_enabled", "quantity_min", "quantity_max"),
                    (amount, "amount_enabled", "amount_min", "amount_max"),
                ):
                    if (
                        value and limit and limit[enabled_key]
                        and limit[min_key] is not None and limit[max_key] is not None
                    ):
                        try:
                            numeric_value = float(value)
                        except ValueError:
                            continue
                        if numeric_value < limit[min_key] or numeric_value > limit[max_key]:
                            control_warning_count += 1
                db.execute(
                    """
                    INSERT INTO energy_entries (
                        report_month, user_id, energy_type_id, area_id, quantity, amount, updated_at
                    )
                    VALUES (?, ?, ?, ?, NULLIF(?, ''), NULLIF(?, ''), CURRENT_TIMESTAMP)
                    ON CONFLICT(report_month, user_id, energy_type_id, area_id)
                    DO UPDATE SET
                        quantity = excluded.quantity,
                        amount = excluded.amount,
                        updated_at = CURRENT_TIMESTAMP
                    """,
                    (report_month, user_id, energy_type["id"], area["id"], quantity, amount),
                )
        db.execute(
            """
            INSERT INTO report_submissions (report_month, user_id)
            VALUES (?, ?)
            ON CONFLICT(report_month, user_id)
            DO UPDATE SET
                submitted_at = CURRENT_TIMESTAMP,
                status = 'submitted'
            """,
            (report_month, user_id),
        )
        db.commit()
        log_action("submit_report", f"Submitted data for {report_month} (Warnings: {control_warning_count})")
        flash(f"{report_month} 能源資料已送出。若需修正，可再次編輯後重新送出。", "success")
        if control_warning_count:
            flash(f"{control_warning_count} 筆資料超出管制範圍，仍可送出。", "warning")
        return redirect(url_for("fill_report", report_month=report_month))

    entries = {
        (row["energy_type_id"], row["area_id"]): row
        for row in db.execute(
            """
            SELECT * FROM energy_entries
            WHERE report_month = ? AND user_id = ?
            """,
            (report_month, user_id),
        ).fetchall()
    }
    return render_template(
        "fill_report.html",
        report_month=report_month,
        max_report_month=max_report_month,
        energy_types=energy_types,
        areas=areas,
        entries=entries,
        control_limits=control_limits,
        is_submitted=is_submitted,
        submitted_at=submission["submitted_at"] if submission else None,
    )


def report_rows(report_month: str, report_year: str, start_month: str = "", end_month: str = "") -> list[sqlite3.Row]:
    filters = []
    params = []
    if start_month and end_month:
        filters.append("ee.report_month BETWEEN ? AND ?")
        params.extend([start_month, end_month])
    elif report_month:
        filters.append("ee.report_month = ?")
        params.append(report_month)
    elif report_year:
        filters.append("substr(ee.report_month, 1, 4) = ?")
        params.append(report_year)
    where_sql = f"WHERE {' AND '.join(filters)}" if filters else ""
    return get_db().execute(
        f"""
        SELECT
            ee.report_month,
            u.display_name,
            rs.submitted_at,
            et.name AS energy_name,
            et.unit,
            a.name AS area_name,
            ee.quantity,
            ee.amount,
            ee.updated_at
        FROM energy_entries ee
        JOIN users u ON u.id = ee.user_id
        JOIN energy_types et ON et.id = ee.energy_type_id
        JOIN areas a ON a.id = ee.area_id
        LEFT JOIN report_submissions rs
            ON rs.report_month = ee.report_month
            AND rs.user_id = ee.user_id
        {where_sql}
        ORDER BY ee.report_month DESC, u.display_name, et.display_order, a.display_order
        """,
        params,
    ).fetchall()


def report_filters() -> tuple[str, str, str, str]:
    report_month = request.args.get("report_month", "").strip()
    report_year = request.args.get("report_year", "").strip()
    start_month = request.args.get("start_month", "").strip()
    end_month = request.args.get("end_month", "").strip()
    mode = request.args.get("mode", "").strip()
    if not mode:
        # Old links with one unambiguous condition remain usable.
        choices = [name for name, present in (
            ("month", bool(report_month)), ("year", bool(report_year)),
            ("range", bool(start_month or end_month)),
        ) if present]
        if len(choices) > 1:
            raise ValueError("查詢條件混用了不同方式，請先選擇一種查詢方式。")
        mode = choices[0] if choices else "all"
    if mode == "range":
        if not start_month or not end_month:
            raise ValueError("請同時填寫起始月份與結束月份。")
        if not valid_report_month(start_month) or not valid_report_month(end_month):
            raise ValueError("月份區間格式不正確。")
        if start_month > end_month:
            raise ValueError("起始月份不可晚於結束月份。")
        return "", "", start_month, end_month
    if mode == "month":
        if not report_month or not valid_report_month(report_month):
            raise ValueError("請選擇有效的查詢月份。")
        return report_month, "", "", ""
    if mode == "year":
        if not re.fullmatch(r"[0-9]{4}", report_year) or not 1 <= int(report_year) <= 9999:
            raise ValueError("請選擇有效的查詢年份。")
        return "", report_year, "", ""
    if mode == "all":
        return "", "", "", ""
    raise ValueError("查詢方式不正確，請重新選擇。")


@app.route("/reports", methods=["GET"])
@app.route("/admin/reports", methods=["GET"])
def admin_reports():
    guard = report_access_required()
    if guard:
        return guard

    available_months = [row[0] for row in get_db().execute(
        "SELECT DISTINCT report_month FROM energy_entries ORDER BY report_month DESC"
    ).fetchall()]
    years = sorted({month[:4] for month in available_months} | {latest_report_month()[:4]}, reverse=True)
    filter_error = False
    try:
        report_month, report_year, start_month, end_month = report_filters()
    except ValueError as exc:
        flash(str(exc), "danger")
        filter_error = True
        report_month, report_year, start_month, end_month = (
            request.args.get(key, "").strip() for key in
            ("report_month", "report_year", "start_month", "end_month")
        )
    mode = request.args.get("mode") or (
        "range" if start_month or end_month else "month" if report_month else "year" if report_year else "all"
    )
    if mode not in {"month", "range", "year", "all"}:
        mode = "month"
    if re.fullmatch(r"[0-9]{4}", report_year) and report_year not in years:
        years = sorted([*years, report_year], reverse=True)
    return render_template(
        "admin_reports.html",
        mode=mode,
        filter_error=filter_error,
        default_month=latest_report_month(),
        available_months=available_months,
        report_month=report_month,
        report_year=report_year,
        start_month=start_month,
        end_month=end_month,
        years=years,
        rows=[] if filter_error else report_rows(report_month, report_year, start_month, end_month),
    ), 400 if filter_error else 200


@app.route("/reports/export", methods=["GET"])
def export_reports():
    guard = report_access_required()
    if guard:
        return guard

    try:
        report_month, report_year, start_month, end_month = report_filters()
    except ValueError:
        return admin_reports()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = translate("能源填報資料")
    sheet.append([
        translate("月份"), translate("填寫員"), translate("能源種類"),
        translate("能源單位"), translate("區域"), translate("數量"),
        translate("金額(未稅)"), translate("送出時間"), translate("更新時間"),
    ])
    for row in report_rows(report_month, report_year, start_month, end_month):
        sheet.append([
            row["report_month"], row["display_name"], row["energy_name"], row["unit"],
            row["area_name"], row["quantity"], row["amount"],
            row["submitted_at"], row["updated_at"],
        ])
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(
            max(len(str(cell.value or "")) for cell in column) + 2,
            40,
        )
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    suffix = f"{start_month}_to_{end_month}" if start_month else report_month or report_year or "all"
    return send_file(
        output,
        as_attachment=True,
        download_name=f"energy-report-{suffix}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


if __name__ == "__main__":
    app.run(debug=True, use_reloader=False)
