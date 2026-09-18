from __future__ import annotations

import hashlib
import json
import re
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from contextlib import contextmanager
from datetime import date
from io import BytesIO
from pathlib import Path
from openpyxl import load_workbook

import app as application


class EmsPointFeatureTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.original_db_path = application.DB_PATH
        self.original_seed_path = application.POINT_SEED_PATH
        self.original_manifest_path = application.POINT_MANIFEST_PATH
        application.DB_PATH = Path(self.temp_dir.name) / "app.db"
        application._INITIALIZED_DB_PATH = None
        application.app.config.update(TESTING=True, SECRET_KEY="test-secret")
        self.client = application.app.test_client()
        self.client.get("/login")

    def tearDown(self) -> None:
        application.DB_PATH = self.original_db_path
        application.POINT_SEED_PATH = self.original_seed_path
        application.POINT_MANIFEST_PATH = self.original_manifest_path
        application._INITIALIZED_DB_PATH = None
        self.temp_dir.cleanup()

    @contextmanager
    def database(self):
        connection = sqlite3.connect(application.DB_PATH)
        connection.row_factory = sqlite3.Row
        try:
            yield connection
        finally:
            connection.close()

    def login_admin(self) -> None:
        response = self.client.post(
            "/login",
            data={"username": "admin", "password": "1234"},
        )
        self.assertEqual(response.status_code, 302)

    def test_package_import_and_energy_mount_in_isolated_process(self) -> None:
        script = r'''
import importlib
import os
import sys
import types
from werkzeug.middleware.dispatcher import DispatcherMiddleware
from werkzeug.test import Client
from werkzeug.wrappers import Response

package = types.ModuleType("energy")
package.__path__ = [sys.argv[1]]
sys.modules["energy"] = package
os.environ["ENERGY_DATABASE_PATH"] = sys.argv[2]
module = importlib.import_module("energy.app")
assert module.StaticUiTranslation.__module__ == "energy.localization"
client = Client(DispatcherMiddleware(Response("Not Found", status=404), {"/energy": module.app}), Response)
response = client.get("/energy")
if response.status_code == 308:
    assert response.location.endswith("/energy/"), response.location
    response = client.get(response.location)
assert response.status_code == 302, response.status_code
assert response.location == "/energy/login", response.location
assert client.get(response.location).status_code == 200
client.post("/energy/login", data={"username": "admin", "password": "1234"})
response = client.get("/energy/admin/ems-points")
assert response.status_code == 200
assert '/energy/admin/ems-points/bulk' in response.get_data(as_text=True)
response = client.post("/energy/language/vi", data={"next": "/energy/admin/ems-points"})
assert response.location == "/energy/admin/ems-points"
response = client.get(response.location)
assert response.status_code == 200
assert 'Cấu hình điểm tiện ích EMS' in response.get_data(as_text=True)
assert client.get("/energy/reports?start_month=2025-10&end_month=2026-03").status_code == 200
'''
        result = subprocess.run(
            [sys.executable, "-I", "-c", script, str(application.BASE_DIR), str(Path(self.temp_dir.name) / "mounted.db")],
            capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=30,
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    def test_fill_defaults_to_previous_month_and_rejects_current_month(self) -> None:
        self.assertEqual(application.latest_report_month(date(2026, 9, 18)), "2026-08")
        self.assertEqual(application.latest_report_month(date(2026, 1, 5)), "2025-12")

        with self.database() as db:
            db.execute(
                "INSERT INTO users (username, password, display_name, role) VALUES (?, ?, ?, ?)",
                ("filler-test", "test-password", "Filler Test", "filler"),
            )
            db.commit()
        login = self.client.post(
            "/login",
            data={"username": "filler-test", "password": "test-password"},
        )
        self.assertEqual(login.status_code, 302)

        expected_month = application.latest_report_month()
        page = self.client.get("/fill")
        self.assertEqual(page.status_code, 200)
        self.assertIn(f'value="{expected_month}"'.encode(), page.data)
        self.assertIn(f'max="{expected_month}"'.encode(), page.data)

        rejected = self.client.post(
            "/fill",
            data={"report_month": application.current_report_month()},
        )
        self.assertEqual(rejected.status_code, 302)
        self.assertIn(f"report_month={expected_month}", rejected.location)
        with self.database() as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM report_submissions").fetchone()[0], 0)

    def test_initial_catalog_import_is_complete_and_idempotent(self) -> None:
        with self.database() as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM device_data_points").fetchone()[0], 388)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM ems_point_configs").fetchone()[0], 388)
            self.assertEqual(
                db.execute("SELECT COUNT(*) FROM ems_point_configs WHERE monitor_enabled = 1").fetchone()[0],
                388,
            )
            versions_before = db.execute("SELECT COUNT(*) FROM point_catalog_versions").fetchone()[0]
        application._INITIALIZED_DB_PATH = None
        self.client.get("/login")
        with self.database() as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM device_data_points").fetchone()[0], 388)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM point_catalog_versions").fetchone()[0], versions_before)

    def test_month_range_query_and_export_for_viewer_and_admin(self) -> None:
        months = ["2025-09", "2025-10", "2025-12", "2026-03", "2026-04", "2026-08"]
        with self.database() as db:
            db.execute(
                "INSERT INTO users (username, password, display_name, role, viewer) VALUES (?, ?, ?, ?, ?)",
                ("range-viewer", "test-password", "Range Viewer", "filler", 1),
            )
            user_id = db.execute("SELECT id FROM users WHERE username = 'admin'").fetchone()[0]
            energy_id = db.execute("SELECT id FROM energy_types LIMIT 1").fetchone()[0]
            area_id = db.execute("SELECT id FROM areas LIMIT 1").fetchone()[0]
            db.executemany(
                "INSERT INTO energy_entries (report_month, user_id, energy_type_id, area_id, quantity) VALUES (?, ?, ?, ?, ?)",
                [(month, user_id, energy_id, area_id, 10) for month in months],
            )
            db.commit()

        for username, password in [("range-viewer", "test-password"), ("admin", "1234")]:
            self.client.post("/login", data={"username": username, "password": password})
            for filters, expected in [
                ({"start_month": "2025-10", "end_month": "2026-03"}, months[1:4]),
                ({"start_month": "2025-10", "end_month": "2025-10"}, ["2025-10"]),
                ({"mode": "range", "start_month": "2025-10", "end_month": "2026-03", "report_month": "2025-09", "report_year": "2025"}, months[1:4]),
                ({"mode": "month", "report_month": "2026-08", "start_month": "2025-10", "end_month": "2026-03", "report_year": "2025"}, ["2026-08"]),
                ({"mode": "year", "report_year": "2025", "report_month": "2026-08"}, months[:3]),
                ({"mode": "all", "report_month": "2026-08"}, months),
                ({"report_month": "2026-03"}, ["2026-03"]),
                ({"report_year": "2025"}, months[:3]),
                ({}, months),
                ({"start_month": "2024-01", "end_month": "2024-12"}, []),
            ]:
                with self.subTest(username=username, filters=filters):
                    page = self.client.get("/reports", query_string=filters)
                    self.assertEqual(page.status_code, 200)
                    for month in months:
                        self.assertEqual(f"<td>{month}</td>".encode() in page.data, month in expected)
                    export = self.client.get("/reports/export", query_string=filters)
                    self.assertEqual(export.status_code, 200)
                    workbook = load_workbook(BytesIO(export.data))
                    actual = [row[0] for row in workbook.active.iter_rows(min_row=2, values_only=True)]
                    self.assertEqual(actual, sorted(expected, reverse=True))
                    workbook.close()
            for filters in [
                {"start_month": "2025-10", "end_month": "2026-03", "report_month": "2026-08"},
                {"mode": "month"},
                {"mode": "month", "report_month": "2026-13"},
                {"mode": "year", "report_year": "bad"},
                {"mode": "invalid"},
                {"start_month": "2025-10"},
                {"end_month": "2026-03"},
                {"start_month": "2026-03", "end_month": "2025-10"},
                {"start_month": "2025-13", "end_month": "2026-03"},
            ]:
                for endpoint in ["/reports", "/reports/export"]:
                    response = self.client.get(endpoint, query_string=filters)
                    self.assertEqual(response.status_code, 400)
                    self.assertNotIn(b"<td>2025-10</td>", response.data)
                    self.assertEqual(response.mimetype, "text/html")

    def test_report_forms_show_only_the_active_query_fields(self) -> None:
        self.login_admin()
        for filters, expected in [
            ({"mode": "month", "report_month": "2026-08"}, {"report_month"}),
            ({"mode": "range", "start_month": "2025-10", "end_month": "2026-08"}, {"start_month", "end_month"}),
            ({"mode": "year", "report_year": "2025"}, {"report_year"}),
            ({"mode": "all"}, set()),
        ]:
            page = self.client.get("/reports", query_string=filters).get_data(as_text=True)
            for field in ["report_month", "start_month", "end_month", "report_year"]:
                self.assertEqual(f'id="{field}"' in page, field in expected)
            self.assertIn("目前結果：", page)
            self.assertIn("已有填報紀錄的月份", page)
        page = self.client.get("/reports?mode=range&start_month=2026-08").get_data(as_text=True)
        self.assertIn('value="2026-08"', page)
        self.assertIn("查詢尚未執行", page)
        self.assertNotIn("匯出目前結果</a>", page)

    def test_ems_export_uses_combined_filters_across_pages(self) -> None:
        self.login_admin()
        with self.database() as db:
            ids = [row[0] for row in db.execute("SELECT id FROM device_data_points ORDER BY external_id LIMIT 30")]
            category_id = db.execute("SELECT id FROM ems_energy_categories LIMIT 1").fetchone()[0]
            db.executemany("UPDATE device_data_points SET description = 'test combined search' WHERE id = ?", [(value,) for value in ids])
            db.executemany("UPDATE ems_point_configs SET monitor_enabled = 0, energy_category_id = ? WHERE point_id = ?", [(category_id, value) for value in ids])
            db.commit()
        filters = {"q": "test combined search", "monitor": "off", "category_id": str(category_id), "per_page": "25"}
        page = self.client.get("/admin/ems-points", query_string=filters).get_data(as_text=True)
        self.assertIn("目前結果：30 筆", page)
        self.assertIn("test combined search", page)
        self.assertEqual(page.count('class="edit-point"'), 25)
        for monitor, expected in [("off", 30), ("on", 0)]:
            export = self.client.get("/admin/ems-points/export", query_string={**filters, "monitor": monitor})
            workbook = load_workbook(BytesIO(export.data))
            self.assertEqual(workbook.active.max_row, expected + 1)
            workbook.close()

    def test_bilingual_ui_preserves_user_data_and_submitted_values(self) -> None:
        self.login_admin()
        original_text = '上限 設備／系統 <script>alert(1)</script>'
        with self.database() as db:
            point_id = db.execute("SELECT id FROM device_data_points ORDER BY external_id LIMIT 1").fetchone()[0]
            db.execute("UPDATE device_data_points SET description = ? WHERE id = ?", (original_text, point_id))
            db.commit()
        self.client.post("/language/vi", data={"next": "/admin/ems-points"})
        for endpoint in ["/admin", "/admin/users", "/admin/settings", "/admin/control-limits", "/reports", "/admin/ems-points"]:
            self.assertEqual(self.client.get(endpoint).status_code, 200)
        page = self.client.get("/admin/ems-points").get_data(as_text=True)
        self.assertIn("Cấu hình điểm tiện ích EMS", page)
        self.assertIn("Giá trị đánh giá cảnh báo", page)
        self.assertIn("Số lần bất thường liên tiếp", page)
        self.assertIn('<option value="輸入">Đầu vào</option>', page)
        self.assertIn('上限 設備／系統 &lt;script&gt;alert(1)&lt;/script&gt;', page)
        self.assertNotIn(original_text, page)
        response = self.client.post(f"/admin/ems-points/{point_id}", data={
            "flow_direction": "輸入", "display_name": original_text,
            "equipment_name": "設備／系統", "notes": "上限", "alarm_consecutive_samples": "2",
        }, follow_redirects=True)
        self.assertIn("đã lưu cấu hình EMS", response.get_data(as_text=True))
        with self.database() as db:
            row = db.execute("SELECT * FROM ems_point_configs WHERE point_id = ?", (point_id,)).fetchone()
            self.assertEqual(row["flow_direction"], "輸入")
            self.assertEqual(row["display_name"], original_text)
            self.assertEqual(row["notes"], "上限")
        export = self.client.get("/admin/ems-points/export")
        workbook = load_workbook(BytesIO(export.data))
        self.assertEqual(workbook.active.cell(1, 8).value, "Giới hạn dưới")
        self.assertEqual(workbook.active.cell(2, 5).value, original_text)
        self.assertEqual(workbook.active.cell(2, 18).value, "輸入")
        workbook.close()
        response = self.client.post(f"/admin/ems-points/{point_id}", data={"alarm_consecutive_samples": "bad"}, follow_redirects=True)
        self.assertIn("phải là số nguyên từ 1 đến 168", response.get_data(as_text=True))
        self.client.post("/language/zh-TW")
        self.assertIn("EMS 公用點位設定", self.client.get("/admin/ems-points").get_data(as_text=True))
        # Template caching must not freeze the first request's language.
        self.client.post("/language/vi")
        self.assertIn("Cấu hình điểm tiện ích EMS", self.client.get("/admin/ems-points").get_data(as_text=True))

    def test_all_static_template_ui_has_vietnamese_translation(self) -> None:
        with application.app.test_request_context():
            application.session["language"] = "vi"
            for path in (application.BASE_DIR / "templates").glob("*.html"):
                for line, kind, value in application.app.jinja_env.lex(path.read_text(encoding="utf-8")):
                    if kind == "data":
                        translated = application.translate_ui_text(value)
                        self.assertIsNone(re.search(r"[\u3400-\u9fff]", translated), f"{path.name}:{line}: {translated}")

    def test_admin_can_view_and_update_a_point(self) -> None:
        self.login_admin()
        page = self.client.get("/admin/ems-points")
        self.assertEqual(page.status_code, 200)
        self.assertIn(b"FQI_U704A.TOTALIZERA.PV", page.data)
        with self.database() as db:
            point_id = db.execute(
                "SELECT id FROM device_data_points WHERE data_source_id = 1 AND external_id = 2"
            ).fetchone()[0]
            category_id = db.execute(
                "SELECT id FROM ems_energy_categories WHERE name = '壓縮空氣'"
            ).fetchone()[0]
        response = self.client.post(
            f"/admin/ems-points/{point_id}",
            data={
                "monitor_enabled": "on",
                "lower_limit": "10.5",
                "upper_limit": "20.5",
                "energy_category_id": str(category_id),
                "alarm_value_mode": "interval_delta",
                "alarm_consecutive_samples": "2",
                "alarm_deadband": "0.5",
                "alarm_severity": "warning",
            },
        )
        self.assertEqual(response.status_code, 302)
        with self.database() as db:
            config = db.execute(
                "SELECT * FROM ems_point_configs WHERE point_id = ?", (point_id,)
            ).fetchone()
            self.assertEqual(config["lower_limit"], 10.5)
            self.assertEqual(config["upper_limit"], 20.5)
            self.assertEqual(config["energy_category_id"], category_id)
            self.assertEqual(config["alarm_value_mode"], "interval_delta")
            self.assertEqual(config["alarm_consecutive_samples"], 2)

    def test_invalid_limits_are_rejected(self) -> None:
        self.login_admin()
        with self.database() as db:
            point_id = db.execute(
                "SELECT id FROM device_data_points WHERE data_source_id = 1 AND external_id = 2"
            ).fetchone()[0]
        self.client.post(
            f"/admin/ems-points/{point_id}",
            data={
                "lower_limit": "100",
                "upper_limit": "10",
                "alarm_value_mode": "raw",
                "alarm_consecutive_samples": "1",
                "alarm_deadband": "0",
                "alarm_severity": "warning",
            },
        )
        with self.database() as db:
            config = db.execute(
                "SELECT lower_limit, upper_limit FROM ems_point_configs WHERE point_id = ?",
                (point_id,),
            ).fetchone()
            self.assertIsNone(config["lower_limit"])
            self.assertIsNone(config["upper_limit"])

    def test_bulk_monitor_update_and_export(self) -> None:
        self.login_admin()
        with self.database() as db:
            point_ids = [
                row[0]
                for row in db.execute(
                    "SELECT id FROM device_data_points WHERE data_source_id = 1 ORDER BY external_id LIMIT 2"
                ).fetchall()
            ]
        response = self.client.post(
            "/admin/ems-points/bulk",
            data={"point_ids": [str(value) for value in point_ids], "action": "monitor_off"},
        )
        self.assertEqual(response.status_code, 302)
        with self.database() as db:
            placeholders = ",".join("?" for _ in point_ids)
            monitored = db.execute(
                f"SELECT SUM(monitor_enabled) FROM ems_point_configs WHERE point_id IN ({placeholders})",
                point_ids,
            ).fetchone()[0]
            self.assertEqual(monitored, 0)
        export = self.client.get("/admin/ems-points/export")
        self.assertEqual(export.status_code, 200)
        self.assertEqual(
            export.mimetype,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_new_catalog_metadata_does_not_overwrite_ems_config(self) -> None:
        with self.database() as db:
            point_id = db.execute(
                "SELECT id FROM device_data_points WHERE data_source_id = 1 AND external_id = 2"
            ).fetchone()[0]
            db.execute(
                "UPDATE ems_point_configs SET lower_limit = 12, upper_limit = 34 WHERE point_id = ?",
                (point_id,),
            )
            db.commit()

        catalog = json.loads(self.original_seed_path.read_text(encoding="utf-8"))
        catalog["points"][0]["description"] = "Changed by source"
        seed_path = Path(self.temp_dir.name) / "utility_points.json"
        manifest_path = Path(self.temp_dir.name) / "utility_points_manifest.json"
        seed_bytes = (json.dumps(catalog, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
        seed_path.write_bytes(seed_bytes)
        manifest_path.write_text(
            json.dumps(
                {
                    "schema_version": 1,
                    "data_source_id": 1,
                    "version": "test.2",
                    "row_count": len(catalog["points"]),
                    "sha256": hashlib.sha256(seed_bytes).hexdigest(),
                }
            ),
            encoding="utf-8",
        )
        application.POINT_SEED_PATH = seed_path
        application.POINT_MANIFEST_PATH = manifest_path
        with application.app.app_context():
            result = application.sync_bundled_device_points(application.get_db())
        self.assertEqual(result["status"], "updated")
        with self.database() as db:
            point = db.execute("SELECT description FROM device_data_points WHERE id = ?", (point_id,)).fetchone()
            config = db.execute(
                "SELECT lower_limit, upper_limit FROM ems_point_configs WHERE point_id = ?", (point_id,)
            ).fetchone()
            self.assertEqual(point["description"], "Changed by source")
            self.assertEqual(config["lower_limit"], 12)
            self.assertEqual(config["upper_limit"], 34)


if __name__ == "__main__":
    unittest.main()
